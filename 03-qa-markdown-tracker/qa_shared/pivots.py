"""Native Excel PivotTables, with Python-calculated opening views and saved caches.

QA uses one row per audit. Markdown rates use one row per audit/markdown,
including blank unknown flags; these rows are never used to average QA.
"""
from collections import defaultdict
from openpyxl.pivot.cache import CacheDefinition, CacheSource, WorksheetSource, CacheField, SharedItems
from openpyxl.pivot.record import Record, RecordList
from openpyxl.pivot.fields import Index, Text, Number, Missing
from openpyxl.pivot.table import (TableDefinition, Location, PivotField, FieldItem, RowColField,
                                  RowColItem, PageField, DataField)
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from .core import CATALOG


def native_cache(sheet, first_col, headers, rows):
    """A populated worksheet source and shared-item records Excel can filter."""
    for j, title in enumerate(headers, first_col):
        sheet.cell(1, j, title)
    for i, row in enumerate(rows, 2):
        for j, value in enumerate(row, first_col):
            cell = sheet.cell(i, j, value)
            if isinstance(value, str):
                cell.data_type = 's'  # Literal names/IDs, never formulas.
    fields, indices = [], []
    for col, title in enumerate(headers):
        unique = list(dict.fromkeys(row[col] for row in rows))
        lookup = {v: i for i, v in enumerate(unique)}
        indices.append(lookup)
        numeric = any(isinstance(v, (int, float)) for v in unique)
        strings = any(isinstance(v, str) for v in unique)
        items = [Missing() if v is None else Number(v=v) if isinstance(v, (int, float)) else Text(v=str(v)) for v in unique]
        fields.append(CacheField(name=title, sharedItems=SharedItems(_fields=items,
            containsBlank=None in unique, containsNumber=numeric, containsString=strings,
            containsNonDate=True, containsMixedTypes=numeric and strings)))
    cache = CacheDefinition(cacheSource=CacheSource(type='worksheet', worksheetSource=WorksheetSource(
        sheet=sheet.title, ref=f'{get_column_letter(first_col)}1:{get_column_letter(first_col+len(headers)-1)}{len(rows)+1}')),
        cacheFields=fields, recordCount=len(rows), saveData=True, refreshOnLoad=False,
        enableRefresh=True, createdVersion=6, refreshedVersion=6, minRefreshableVersion=3)
    cache.records = RecordList(r=[Record(_fields=[Index(v=indices[j][value]) for j, value in enumerate(row)]) for row in rows])
    return cache, indices


def aggregate(rows, measures):
    values = []
    for field, operation, _, _ in measures:
        observed = [row[field] for row in rows if row[field] is not None]
        if operation == 'count':
            values.append(len(observed))
        elif operation == 'sum':
            values.append(sum(observed) if observed else None)
        else:
            values.append(sum(observed) / len(observed) if observed else None)
    return values


def add_pivot(ws, anchor, name, cache, indices, rows, label_field, page_fields, measures):
    """Single label axis plus measures: simple browsing and editable field lists."""
    top, col = anchor
    buckets = defaultdict(list)
    for row in rows:
        buckets[row[label_field]].append(row)
    labels = sorted(buckets)
    header = top + len(page_fields) + 1
    last = header + len(labels) + 1
    for i, field in enumerate(page_fields, top):
        ws.cell(i, col, cache.cacheFields[field].name)
        ws.cell(i, col + 1, '(All)')
        ws.cell(i, col + 1).fill = PatternFill('solid', fgColor='FFF0CD')
    headers = [cache.cacheFields[label_field].name, *(m[2] for m in measures)]
    for j, title in enumerate(headers, col):
        cell = ws.cell(header, j, title)
        cell.font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='233E62')
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    ws.row_dimensions[header].height = 30
    for i, label in enumerate(labels, header + 1):
        ws.cell(i, col, label).data_type = 's'
        for j, value in enumerate(aggregate(buckets[label], measures), col + 1):
            ws.cell(i, j, value)
        for j in range(col, col + len(headers)):
            ws.cell(i, j).font = Font(name='Arial', size=10, color='182B42')
            ws.cell(i, j).alignment = Alignment(vertical='center',wrap_text=j==col)
            if (i-header) % 2 == 0:
                ws.cell(i, j).fill = PatternFill('solid', fgColor='F3F6FA')
        ws.row_dimensions[i].height = 38 if len(str(label))>46 else 24
    ws.cell(last, col, 'Grand Total')
    for j, value in enumerate(aggregate(rows, measures), col + 1):
        ws.cell(last, j, value)
    for j in range(col, col + len(headers)):
        ws.cell(last, j).font = Font(name='Arial', size=10, bold=True, color='182B42')
        ws.cell(last, j).fill = PatternFill('solid', fgColor='EDF2F7')
    for j, (_, _, _, fmt) in enumerate(measures, col + 1):
        for i in range(header + 1, last + 1):
            ws.cell(i, j).number_format = '0.0%' if fmt == 10 else '#,##0'
    pivot_fields = []
    for j, field in enumerate(cache.cacheFields):
        active = j == label_field or j in page_fields
        items = [FieldItem(x=i) for i in range(len(indices[j]))] if active else []
        pivot_fields.append(PivotField(axis='axisRow' if j == label_field else 'axisPage' if j in page_fields else None,
            items=items, dataField=any(m[0] == j for m in measures), defaultSubtotal=False,
            compact=False, outline=False, showAll=False))
    caches = []
    for sheet in ws.parent.worksheets:
        for existing in sheet._pivots:
            if not any(existing.cache is c for c in caches):
                caches.append(existing.cache)
    cache_id = next((i+1 for i,c in enumerate(caches) if c is cache),len(caches)+1)
    pivot = TableDefinition(name=name, cacheId=cache_id, dataCaption='Values',
        location=Location(ref=f'{get_column_letter(col)}{header}:{get_column_letter(col+len(headers)-1)}{last}',
            firstHeaderRow=1, firstDataRow=1, firstDataCol=1, rowPageCount=len(page_fields), colPageCount=1),
        pivotFields=pivot_fields, rowFields=[RowColField(x=label_field)],
        rowItems=[RowColItem(x=[Index(v=indices[label_field][label])]) for label in labels] + [RowColItem(t='grand', x=[Index(v=0)])],
        colFields=[RowColField(x=-2)],
        colItems=[RowColItem(i=i, x=[Index(v=i)]) for i in range(len(measures))],
        pageFields=[PageField(fld=i, hier=-1) for i in page_fields],
        dataFields=[DataField(fld=m[0], subtotal=m[1], name=m[2], numFmtId=m[3]) for m in measures],
        compact=False, compactData=False, outline=False, outlineData=False,
        rowGrandTotals=False, colGrandTotals=True, updatedVersion=6, createdVersion=6,
        minRefreshableVersion=3, preserveFormatting=True, useAutoFormatting=False)
    pivot.cache = cache
    ws.add_pivot(pivot)
    return last


def add_pivot_page(wb, report, setup_sheet):
    ws = wb.create_sheet('Pivot Tables', 1)
    ws.sheet_properties.tabColor = 'C28A35'
    setup_sheet(ws, 'Browse QA with PivotTables', report)
    ws['B6'] = 'Choose the yellow filter dropdowns. Double-click a number to open its source rows in Excel.'
    ws['B6'].font = Font(name='Arial', size=10, color='61738A')
    ws['B7'] = 'QA uses all MTD audits in the chosen team and call type. Markdown rates use the selected analysis dates.'
    ws['B7'].font = Font(name='Arial', size=10, color='61738A')
    for col in ['B','I']:
        ws.column_dimensions[col].width = 48
    for col in ['C','D','E','J','K','L']:
        ws.column_dimensions[col].width = 18
    ws.column_dimensions['F'].width = 3
    ws.column_dimensions['G'].width = 3
    ws.column_dimensions['H'].width = 3
    source = wb.create_sheet('Pivot Source')
    source.sheet_state = 'hidden'
    qa_rows = [[a['audit_id'], a['agent'] + ' (' + a['agent_id'] + ')', a['supervisor'], a['manager'], a['call_type'],
                a['date'], a['score']/100 if a['score'] is not None else None, a['transaction_id']]
               for a in report.get('pivot_qa', report['detail'])]
    qa_headers = ['Audit ID','Agent (Login ID)','Supervisor','Manager','Call type','Audit date','QA score','Transaction ID']
    if qa_rows:
        qa_cache, qa_idx = native_cache(source, 1, qa_headers, qa_rows)
        ws['B9'] = 'QA MTD by supervisor'
        end = add_pivot(ws, (10, 2), 'QATeams', qa_cache, qa_idx, qa_rows, 2, [3,4],
                        [(6,'average','QA MTD',10),(0,'count','Audits',3),(6,'count','Scored audits',3)])
        ws.cell(end+3, 2, 'QA MTD by agent')
        add_pivot(ws, (end+4,2), 'QAAgents', qa_cache, qa_idx, qa_rows, 1, [2,3,4],
                   [(6,'average','QA MTD',10),(0,'count','Audits',3),(6,'count','Scored audits',3)])
    else:
        ws['B9'] = 'No audited QA results are available. Import the QA audit export to populate these pivots.'
        ws['B9'].alignment = Alignment(wrap_text=True)
        ws.row_dimensions[9].height = 43
    names = list(report['filters'].get('markdown') and [report['filters']['markdown']] or CATALOG)
    md_rows = [[a['audit_id'], a['agent']+' ('+a['agent_id']+')',a['supervisor'],a['manager'],a['call_type'],m,
                a['call_type']+' / '+m, int(a['markdowns'][m]) if a['markdowns'].get(m) is not None else None,a['transaction_id']]
               for a in report.get('pivot_markdowns',report['detail']) for m in names]
    right_end = 10
    if md_rows:
        md_cache, md_idx = native_cache(source, 10,
            ['Audit ID','Agent (Login ID)','Supervisor','Manager','Call type','Markdown','Call type / Markdown','Flag (1=Yes, 0=No)','Transaction ID'], md_rows)
        ws['I9'] = 'Markdown frequency by call type'
        right_end = add_pivot(ws,(10,9),'MarkdownFrequency',md_cache,md_idx,md_rows,6,[2,3,1,5,4],
            [(7,'average','Frequency',10),(7,'sum','Flags',3),(7,'count','Evaluated audits',3)])
        ws.cell(right_end+2,9,'Unknown flags are excluded. Markdown rows are not used for QA averages.')
        ws.cell(right_end+2,9).font = Font(name='Arial',size=10,color='61738A')
    sample_rows = [[s['transaction_id'],s['agent']+' ('+s['agent_id']+')',s['review_status'],s['recording'],s['call_date']]
                   for s in report['sample']]
    if sample_rows:
        sample_cache,sample_idx = native_cache(source,20,['Transaction ID','Sample agent (Login ID)','Review status','Recording','Call date'],sample_rows)
        sample_top = max(right_end+5,25)
        ws.cell(sample_top,9,'Supplied call sample by agent (all selected calls)')
        add_pivot(ws,(sample_top+1,9),'SampleCoverage',sample_cache,sample_idx,sample_rows,1,[2,3],[(0,'count','Selected calls',3)])
    ws.freeze_panes = 'C13'
    return ws
