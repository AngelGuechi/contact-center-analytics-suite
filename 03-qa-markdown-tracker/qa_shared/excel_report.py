"""Python owns every calculated value; the workbook is an inspectable snapshot."""
from __future__ import annotations

from math import ceil
from pathlib import Path
import os
import tempfile

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.utils import get_column_letter

from .core import CATALOG, audit_display
from .pivots import add_pivot_page

INK, BLUE, TEAL, MUTED = "182B42", "233E62", "0D746D", "61738A"
FILLS = {"Needs attention": "FDE7E6", "Improving": "DCF4E9", "Stable": "EDF2F7",
         "Low sample": "FFF0CD", "No data": "EDF2F7", "On track": "DCF4E9"}


def text_value(value):
    # User text must never become an executable Excel formula.
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def pct(value):
    return value / 100 if value is not None else "No data"


def table(ws, start, headers, rows, name, widths=None):
    first_row, first_col = start
    for offset, header in enumerate(headers):
        cell = ws.cell(first_row, first_col + offset, header)
        cell.font = Font(name="Arial", size=10, color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[first_row].height = 30
    for offset, values in enumerate(rows, 1):
        height = 25
        for index, value in enumerate(values):
            cell = ws.cell(first_row + offset, first_col + index, text_value(value))
            cell.font = Font(name="Arial", size=10, color=INK)
            cell.alignment = Alignment(vertical="center", wrap_text=True,
                                       horizontal="right" if isinstance(value, (int, float)) else "left")
            if offset % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F3F6FA")
            if isinstance(value, str):
                for status, fill in FILLS.items():
                    if value.startswith(status):
                        cell.fill = PatternFill("solid", fgColor=fill)
                        break
                width = (widths or {}).get(index, 18)
                height = max(height, min(180, ceil(len(value) / max(width - 2, 5)) * 14 + 7))
        ws.row_dimensions[first_row + offset].height = height
    for index, width in (widths or {}).items():
        ws.column_dimensions[get_column_letter(first_col + index)].width = width
    if rows:
        ref = f"{get_column_letter(first_col)}{first_row}:{get_column_letter(first_col + len(headers) - 1)}{first_row + len(rows)}"
        t = Table(displayName=name, ref=ref)
        t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(t)
    return first_row + len(rows)


def setup_sheet(ws, title, report):
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_view.zoomScale = 80
    ws.column_dimensions["A"].width = 3
    for col in range(2, 20):
        ws.column_dimensions[get_column_letter(col)].width = 18
    ws["B2"] = title
    ws["B2"].font = Font(name="Arial", size=16, bold=True, color=INK)
    ws.row_dimensions[2].height = 29
    ws["B3"] = f"{report['mode']} | Report date: {report['report_date']} | Last successful refresh: {report['built_at']}"
    ws["B3"].font = Font(name="Arial", size=10, color=MUTED)
    ws.row_dimensions[3].height = 23
    if report["mode"] == "Synthetic demo":
        ws["B4"] = ("SIMULATED QA RESULTS. Agent names are from the supplied file; scores, teams and coaching are examples."
                    if report.get('demo_names_from_sample') else
                    "SYNTHETIC QA RESULTS. These DEMO transactions are separate from the supplied call-sampling file.")
        ws["B4"].font = Font(name="Arial", size=10, color="8B5C09", bold=True)
    ws["B5"] = f"Markdown analysis: {report['start_date']} to {report['report_date']}. QA MTD always starts {report['month_start']}."
    ws["B5"].font = Font(name="Arial", size=10, color=MUTED)


def link(cell, row, label="View audits"):
    cell.value = label
    cell.hyperlink = f"#'Audit Detail'!B{row}"
    cell.font = Font(name="Arial", size=10, color=TEAL, underline="single")


def write_report(report, target):
    wb = Workbook()
    overview = wb.active
    overview.title = "Overview"
    markdowns = wb.create_sheet("Markdowns")
    detail = wb.create_sheet("Audit Detail")
    overview.sheet_properties.tabColor = BLUE
    markdowns.sheet_properties.tabColor = TEAL
    for ws, title in [(overview, "QA performance"), (markdowns, "Markdown patterns"), (detail, "Audit evidence")]:
        setup_sheet(ws, title, report)

    display = [audit_display(a) for a in report["detail"]]
    headers = list(display[0]) if display else list(audit_display({
        "audit_id": "", "transaction_id": "", "date": "", "agent": "", "agent_id": "", "supervisor": "", "manager": "",
        "call_type": "", "score": None, "markdowns": {}, "evaluator": "", "comment": "", "dispositions": []}))
    headers += list(CATALOG)
    detail["B6"] = "Links jump to the first matching audit. Use table-header filters to review the matching agent, call type, and markdown."
    detail["B6"].font = Font(name="Arial", size=10, color=MUTED)
    values = [[*item.values(), *("Yes" if a["markdowns"].get(m) is True else "No" if a["markdowns"].get(m) is False else "Unknown" for m in CATALOG)]
              for item, a in zip(display, report["detail"])]
    table(detail, (8, 2), headers, values, "AuditsTable", {0: 23, 1: 24, 2: 14, 3: 21, 4: 14,
          5: 14, 6: 16, 7: 18, 8: 12, 9: 36, 10: 26, 11: 22, 12: 68, 13: 55,
          14: 12, 15: 12, 16: 18, 17: 12, 18: 12})
    for row in range(9, 9 + len(values)):
        detail.cell(row, 10).number_format = '0.0"%"'
    detail.freeze_panes = "F9"
    detail.print_title_rows = "8:8"
    audit_rows = {a["audit_id"]: 9 + i for i, a in enumerate(report["detail"])}
    sample_row = 12 + len(values)
    detail.cell(sample_row, 2, "SUPPLIED CALL SAMPLE - separate from synthetic QA results").font = Font(name="Arial", size=12, bold=True, color=INK)
    samples = [[s["transaction_id"], s["agent"], s["call_date"], s["duration"], s["recording"],
                s["sample_category"], s["review_status"]] for s in report["sample"]]
    table(detail, (sample_row + 2, 2), ["Transaction ID", "Sample agent", "Call date", "Duration", "Recording", "Sample category", "Review status"],
          samples, "SelectedCallsTable")
    issue_row = sample_row + len(samples) + 6
    detail.cell(issue_row, 2, "Import issues").font = Font(name="Arial", size=12, bold=True, color=INK)
    table(detail, (issue_row + 1, 2), ["File", "Row", "Issue"],
          [[i["file"], i["row"], i["issue"]] for i in report["issues"]], "ImportIssuesTable")

    overview.merge_cells("B7:M8")
    overview["B7"] = report["summary"] + " This text is rule-based."
    overview["B7"].font = Font(name="Arial", size=11, color=INK)
    overview["B7"].alignment = Alignment(wrap_text=True, vertical="center")
    overview.row_dimensions[7].height = 27
    overview.row_dimensions[8].height = 27
    metrics = [("B10", "QA MTD", pct(report["qa"]), "0.0%"),
               ("E10", "Completed audits", report["audit_count"], "#,##0"),
               ("H10", "Change vs last month", report["change_pp"] if report["change_pp"] is not None else "No data", '0.0" pp"'),
               ("K10", "QA attention alerts", sum("Needs attention" in a["qa_alert"] for a in report["agents"]), "#,##0")]
    for anchor, label, value, fmt in metrics:
        overview[anchor] = label
        overview[anchor].font = Font(name="Arial", size=10, color=MUTED)
        cell = overview.cell(overview[anchor].row + 1, overview[anchor].column, value)
        cell.font = Font(name="Arial", size=16, color=INK, bold=True)
        cell.number_format = fmt
    overview.row_dimensions[11].height = 30
    overview["B13"] = f"Last month: {report['previous_start']} to {report['previous_end']} ({report['previous_count']} audits). Missing/invalid MTD scores: {report['audit_count'] - report['scored_audits']}."
    overview["B13"].font = Font(name="Arial", size=10, color=MUTED)
    overview["B14"] = f"Selected sample calls: {len(report['sample'])}; reviewed: {sum(s['review_status'] == 'Reviewed' for s in report['sample'])}. Import issues: {len(report['issues'])}."
    overview["B14"].font = Font(name="Arial", size=10, color=MUTED)
    overview["B16"] = "Team and manager results"
    overview["B16"].font = Font(name="Arial", size=12, bold=True, color=INK)
    teams = report["rollups"]["supervisor"]
    table(overview, (18, 2), ["Supervisor", "QA MTD", "Audits", "Scored audits"],
          [[g["group"], pct(g["qa"]), g["audits"], g["scored_audits"]] for g in teams], "TeamsTable")
    managers = report["rollups"]["manager"]
    table(overview, (18, 9), ["Manager", "QA MTD", "Audits", "Scored audits"],
          [[g["group"], pct(g["qa"]), g["audits"], g["scored_audits"]] for g in managers], "ManagersTable")
    for idx in range(len(teams)):
        overview.cell(19 + idx, 3).number_format = "0.0%"
    for idx in range(len(managers)):
        overview.cell(19 + idx, 10).number_format = "0.0%"
    agent_header = max(28, 22 + len(teams), 22 + len(managers))
    overview.cell(agent_header - 2, 2, "Agents and coaching focus").font = Font(name="Arial", size=12, bold=True, color=INK)
    agent_values = [[a["agent"], a["supervisor"], pct(a["qa"]), a["audits"], a["qa_alert"],
                     a["focus"]["markdown"] if a["focus"] else "No recorded markdown",
                     *(next((s["status"] for s in a["statuses"] if s["markdown"] == md), "Not selected") for md in CATALOG), "View audits"]
                    for a in report["agents"]]
    table(overview, (agent_header, 2), ["Agent", "Supervisor", "QA MTD", "Audits", "QA alert", "Coaching focus", *CATALOG, "Evidence"],
          agent_values, "AgentsTable", {0: 21, 1: 14, 2: 12, 3: 9, 4: 28, 5: 24,
                                      6: 19, 7: 19, 8: 20, 9: 19, 10: 19, 11: 14})
    for i, a in enumerate(report["agents"], agent_header + 1):
        overview.cell(i, 4).number_format = "0.0%"
        first = next((audit_rows[x] for x in a["audit_ids"] if x in audit_rows), None)
        if first:
            link(overview.cell(i, 13), first)
        else:
            overview.cell(i, 13, "No matching audit")
    overview.freeze_panes = f"D{agent_header + 1}"
    overview.print_title_rows = f"{agent_header}:{agent_header}"

    markdowns["B7"] = "Markdown frequency by call type (percentage of evaluated audits)"
    markdowns["B7"].font = Font(name="Arial", size=12, bold=True, color=INK)
    call_types = sorted({r["call_type"] for r in report["heatmap"]})
    heat = [[ct, *(pct(next((h["rate"] for h in report["heatmap"] if h["call_type"] == ct and h["markdown"] == md), None)) for md in CATALOG)] for ct in call_types]
    table(markdowns, (9, 2), ["Call type", *CATALOG], heat, "HeatmapTable", {0: 22, 1: 15, 2: 15, 3: 21, 4: 15, 5: 15})
    if call_types:
        markdowns.conditional_formatting.add(f"C10:G{9 + len(call_types)}", ColorScaleRule(start_type="num", start_value=0, start_color="E1F1ED",
                 end_type="num", end_value=1, end_color="CB5C54"))
        for row in range(10, 10 + len(call_types)):
            for col in range(3, 8):
                markdowns.cell(row, col).number_format = "0.0%"
    table(markdowns, (9, 10), ["Markdown", "Audit category", "Listed impact"],
          [[m, c["category"], c["impact"] / 100] for m, c in CATALOG.items()], "CatalogTable", {0: 22, 1: 21, 2: 15})
    for row in range(10, 15):
        markdowns.cell(row, 12).number_format = "0%"
    stats_row = max(19, 13 + len(call_types))
    stats = report["markdowns"]
    table(markdowns, (stats_row, 2), ["Markdown", "Evaluated", "Flagged", "Frequency", "Points lost", "Points per audit", "Evidence"],
          [[s["markdown"], s["evaluated"], s["misses"], pct(s["rate"]), s["points_lost"] if s["points_lost"] is not None else "Not confirmed",
            s["points_per_audit"] if s["points_per_audit"] is not None else "Not confirmed", "View audits"] for s in stats], "MarkdownsTable")
    for row, s in enumerate(stats, stats_row + 1):
        markdowns.cell(row, 5).number_format = "0.0%"
        markdowns.cell(row, 7).number_format = "0.00"
        first = next((audit_rows[x] for x in s["audit_ids"] if x in audit_rows), None)
        if first:
            link(markdowns.cell(row, 8), first)
        else:
            markdowns.cell(row, 8, "No matching audit")
    status_row = stats_row + len(stats) + 5
    markdowns.cell(status_row - 2, 2, "Agent status and comparison windows").font = Font(name="Arial", size=12, bold=True, color=INK)
    status_rows = []
    for a in report["agents"]:
        for s in a["statuses"]:
            bounds, previous, recent = s["windows"], s["prior"], s["recent"]
            status_rows.append([a["agent"], s["markdown"], s["status"],
              f"{bounds['prior_start']} to {bounds['prior_end']}" if bounds else "No comparison yet",
              previous["evaluated"] if previous else 0, previous["misses"] if previous else 0,
              f"{bounds['recent_start']} to {bounds['recent_end']}" if bounds else "No comparison yet",
              recent["evaluated"] if recent else 0, recent["misses"] if recent else 0,
              s["change_pp"] if s["change_pp"] is not None else "No data", bounds["basis"] if bounds else "Awaiting audits",
              a["grow"]["options"] if a["focus"] and a["focus"]["markdown"] == s["markdown"] else ""])
    table(markdowns, (status_row, 2), ["Agent", "Markdown", "Status", "Prior dates", "Prior audits", "Prior flags", "Recent dates", "Recent audits", "Recent flags", "Change (pp)", "Comparison", "Coaching suggestion"],
          status_rows, "StatusesTable", {0: 21, 1: 21, 2: 25, 3: 30, 4: 12, 5: 12, 6: 30, 7: 12, 8: 12, 9: 15, 10: 29, 11: 58})
    markdowns.freeze_panes = f"D{status_row + 1}"
    markdowns["B6"] = "Status defaults: 4 audits per window; 10 pp change; QA attention below 78%. Listed impacts are preserved; exported deductions supply actual loss."
    markdowns["B6"].font = Font(name="Arial", size=10, color=MUTED)
    add_pivot_page(wb, report, setup_sheet)
    for ws in (overview, markdowns, detail, wb['Pivot Tables']):
        for col, title in [(2,'Overview'),(5,'Pivot Tables'),(9,'Markdowns'),(12,'Audit Detail')]:
            cell = ws.cell(1,col,title)
            cell.hyperlink = f"#'{title}'!B2"
            cell.font = Font(name='Arial',size=10,color=TEAL,underline='single')
    wb.active = 0
    wb.properties.title = "QA Performance and Markdown Tracker"
    wb.properties.description = "Deterministic Python report. Synthetic demo results are explicitly labelled."
    target = Path(target)
    target.parent.mkdir(parents=True, exist_ok=True)
    handle, temp = tempfile.mkstemp(prefix="qa-", suffix=".xlsx", dir=target.parent)
    os.close(handle)
    try:
        wb.save(temp)
        os.replace(temp, target)
    finally:
        if Path(temp).exists():
            Path(temp).unlink()
    return target
