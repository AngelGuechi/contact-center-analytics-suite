from __future__ import annotations

import calendar
import hashlib
import json
import math
import re
from collections import defaultdict
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from zipfile import ZipFile
import xml.etree.ElementTree as ET
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

CATALOG = {
    "OIR": {"category": "Resolution", "impact": 40},
    "Silence": {"category": "Interaction", "impact": 20},
    "Mailing Address": {"category": "Documentation", "impact": 20},
    "TPS": {"category": "Resolution", "impact": 20},
    "Notes": {"category": "Documentation", "impact": 10},
}
RULES = {"min_audits": 4, "change_pp": 10.0, "high_rate": 20.0, "qa_attention": 78.0}
TIMEZONE = "America/Mexico_City"


class DataError(ValueError):
    """An input cannot be published safely; preserve the previous report."""


def digest(value):
    return hashlib.sha256(json.dumps(value, sort_keys=True, ensure_ascii=False).encode()).hexdigest()


def clean(value):
    return str(value).strip() if value is not None else ""


def key(value):
    return re.sub(r"[^a-z0-9]", "", clean(value).lower())


def day(value):
    if isinstance(value, datetime):
        if value.tzinfo:
            value = value.astimezone(ZoneInfo(TIMEZONE))
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = clean(value)
    try:
        parsed = datetime.fromisoformat(text.replace("Z", "+00:00"))
        return day(parsed)
    except ValueError:
        pass
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%m/%d/%Y"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            continue
    raise DataError("Invalid date; use an Excel date or YYYY-MM-DD.")


def number(value, percent_format=False):
    if value is None or clean(value) == "":
        return None
    if isinstance(value, bool):
        raise ValueError("A boolean is not a score or deduction")
    text = clean(value)
    n = float(text.rstrip("%"))
    if percent_format and not text.endswith("%"):
        n *= 100
    if not math.isfinite(n):
        raise ValueError("Non-finite number")
    return n


def flag(value):
    text = key(value)
    if text in ("yes", "y", "true", "1", "triggered", "fail", "failed"):
        return True
    if text in ("no", "n", "false", "0", "pass", "passed"):
        return False
    if text in ("", "unknown", "na", "notapplicable"):
        return None
    raise ValueError("Use Yes, No, or leave a markdown result blank")


def rows_xlsx(path, required=()):
    """Read the first sheet, preserving percentage cell semantics."""
    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        ws = wb.worksheets[0]
        iterator = ws.iter_rows()
        headers = next(iterator, ())
        names = [key(c.value) for c in headers]
        if len([x for x in names if x]) != len(set(x for x in names if x)):
            raise DataError(f"Duplicate headers in {path.name}")
        missing = [name for name in required if key(name) not in names]
        if missing:
            raise DataError(f"{path.name}: missing columns: {', '.join(missing)}")
        result = []
        for line, cells in enumerate(iterator, 2):
            if not any(c.value is not None for c in cells):
                continue
            row = {names[i]: c.value for i, c in enumerate(cells) if i < len(names) and names[i]}
            row["_percent"] = [names[i] for i, c in enumerate(cells)
                               if i < len(names) and "%" in c.number_format]
            row["_source"] = path.name
            row["_line"] = line
            result.append(row)
        return result
    finally:
        wb.close()


def selected_calls(path):
    """ODS sample import. Preserve the supplied agent names and stable login IDs."""
    if not path or not Path(path).exists():
        return []
    ns = {"t": "urn:oasis:names:tc:opendocument:xmlns:table:1.0",
          "x": "urn:oasis:names:tc:opendocument:xmlns:text:1.0",
          "o": "urn:oasis:names:tc:opendocument:xmlns:office:1.0"}
    with ZipFile(path) as archive:
        root = ET.fromstring(archive.read("content.xml"))
    result = []
    for table in root.findall(".//t:table", ns):
        if table.get("{" + ns["t"] + "}name") != "Selected Calls":
            continue
        headers = None
        for row in table.findall("t:table-row", ns):
            values = []
            for cell in row:
                if cell.tag not in ("{" + ns["t"] + "}table-cell", "{" + ns["t"] + "}covered-table-cell"):
                    continue
                value = " ".join("".join(p.itertext()) for p in cell.findall("x:p", ns)).strip()
                value = value or cell.get("{" + ns["o"] + "}value", "")
                repeat = int(cell.get("{" + ns["t"] + "}number-columns-repeated", "1"))
                values.extend([value] * min(repeat, 32))
            if not any(values):
                continue
            if headers is None:
                headers = [key(v) for v in values]
                continue
            row = dict(zip(headers, values))
            txn, login = clean(row.get("transaction")), clean(row.get("loginid"))
            if not txn:
                continue
            result.append({"transaction_id": txn, "agent_id": login,
                           "agent": clean(row.get("agent")) or login, "call_date": day(row.get("datetime")),
                           "duration": row.get("durationnormalized", ""),
                           "recording": "Voice and screen" if key(row.get("recordingtype")) == "voiceandscreen"
                           else clean(row.get("recordingtype")),
                           "sample_category": row.get("samplecategory", "")})
    ids = [c["transaction_id"] for c in result]
    if len(ids) != len(set(ids)):
        raise DataError("Duplicate transaction IDs in the selected-call sample")
    return result


def load_sources(input_dir, sample_path=None):
    input_dir = Path(input_dir)
    issues, roster, audits, coaching, tracker = [], {}, {}, [], defaultdict(list)
    calls = selected_calls(sample_path)
    sample_names = {s['agent_id']: s['agent'] for s in calls if s['agent']}
    synthetic = (input_dir / 'SYNTHETIC-DEMO.json').exists()
    # The user requested source names on the simulated QA rows as well as the sample.
    # Keep synthetic IDs and transactions, so these never become real sampled audits.
    sample_people = {}
    for call in calls:
        sample_people.setdefault(call['agent_id'], call['agent'])
    simulated_names = {f'DEMO{i:03}': name for i, name in enumerate(sample_people.values(), 1)}
    source_names_applied = False

    def issue(row, message):
        issues.append({"file": row.get("_source", ""), "row": row.get("_line", ""), "issue": message})

    roster_path = input_dir / "Roster.xlsx"
    if roster_path.exists():
        for r in rows_xlsx(roster_path, ("Login ID", "Supervisor", "Manager", "Site")):
            login = clean(r.get("loginid"))
            if not login:
                issue(r, "Roster row has no login ID")
                continue
            person = {"agent": clean(r.get("agent")) or clean(r.get("agentname")),
                      **{f: clean(r.get(f)) or "Unassigned" for f in ("supervisor", "manager", "site")}}
            if synthetic and login in simulated_names and re.fullmatch(r'Demo agent \d+', person['agent']):
                person['agent'] = simulated_names[login]
                source_names_applied = True
            if login in roster and roster[login] != person:
                raise DataError("Conflicting team assignments for one login ID")
            roster[login] = person

    required = ("Audit ID", "Audit Date", "Transaction ID", "Login ID", "Final QA Score", "Call Type")
    files = sorted(input_dir.glob("Audits*.xlsx"))
    for path in files:
        for r in rows_xlsx(path, required):
            ident, login, txn = (clean(r.get(f)) for f in ("auditid", "loginid", "transactionid"))
            if not ident or not login or not txn:
                issue(r, "Missing audit, login, or transaction ID; audit excluded")
                continue
            try:
                audit_day = day(r.get("auditdate"))
                revision = int(r.get("revision") or 0)
                if revision < 0:
                    raise ValueError("Negative revision")
            except (ValueError, DataError):
                issue(r, "Invalid audit date or revision; audit excluded")
                continue
            try:
                score = number(r.get("finalqascore"), "finalqascore" in r["_percent"])
                if score is not None and not 0 <= score <= 100:
                    raise ValueError("Score outside range")
                if score is None:
                    issue(r, "Missing score; excluded from QA average")
            except (ValueError, TypeError):
                score = None
                issue(r, "Invalid score; excluded from QA average")
            flags, deductions = {}, {}
            for markdown in CATALOG:
                try:
                    flags[markdown] = flag(r.get(key(markdown)))
                except ValueError:
                    flags[markdown] = None
                    issue(r, f"{markdown} result is invalid; excluded from its denominator")
                points_key = key(markdown + " Points Lost")
                try:
                    value = number(r.get(points_key), points_key in r["_percent"])
                    if value is not None and (not 0 <= value <= 100 or flags[markdown] is None
                                              or (flags[markdown] is False and value != 0)):
                        raise ValueError("Deduction does not match the recorded flag")
                    deductions[markdown] = value
                except (ValueError, TypeError):
                    deductions[markdown] = None
                    issue(r, f"{markdown} deduction is invalid; impact remains unconfirmed")
            audit = {"audit_id": ident, "agent_id": login, "transaction_id": txn, "date": audit_day,
                     "source_agent": clean(r.get("agent")) or clean(r.get("agentname")),
                     "score": score, "call_type": clean(r.get("calltype")) or "Unknown",
                     "evaluator": clean(r.get("evaluator")) or "Not supplied",
                     "comment": clean(r.get("comments")), "markdowns": flags,
                     "deductions": deductions, "revision": revision,
                     "source": path.name, "source_row": r["_line"]}
            previous = audits.get(ident)
            if previous:
                meaningful = lambda a: {k: v for k, v in a.items() if k not in ("source", "source_row")}
                if revision == previous["revision"] and meaningful(audit) != meaningful(previous):
                    raise DataError(f"Conflicting audit revision for {ident}; supply a higher Revision")
                if revision <= previous["revision"]:
                    continue
            audits[ident] = audit

    tracker_path = input_dir / "SNOM3.xlsx"
    if tracker_path.exists():
        seen = set()
        for r in rows_xlsx(tracker_path, ("Transaction ID", "User ID", "Markdown")):
            txn, login, md = (clean(r.get(f)) for f in ("transactionid", "userid", "markdown"))
            canonical = next((m for m in CATALOG if key(m) == key(md)), None)
            if not canonical or not txn or not login:
                issue(r, "SNOM3 row has missing IDs or an unknown markdown; detail excluded")
                continue
            detail = {"markdown": canonical, "l1": clean(r.get("l1")), "l2": clean(r.get("l2")),
                      "screen_recording": clean(r.get("screenrecording"))}
            fingerprint = (txn, login, digest(detail))
            if fingerprint not in seen:
                tracker[(txn, login)].append(detail)
                seen.add(fingerprint)
        audit_keys = {(a["transaction_id"], a["agent_id"]) for a in audits.values()}
        for txn, login in tracker:
            if (txn, login) not in audit_keys:
                issues.append({"file": "SNOM3.xlsx", "row": "", "issue": "Tracker transaction/login has no matching audit"})

    coach_path = input_dir / "Coaching.xlsx"
    if coach_path.exists():
        for r in rows_xlsx(coach_path, ("Login ID", "Coaching Date", "Markdown")):
            try:
                md = next((m for m in CATALOG if key(m) == key(r.get("markdown"))), None)
                if not md or not clean(r.get("loginid")):
                    raise DataError("Invalid coaching ID or markdown")
                entry = {"agent_id": clean(r["loginid"]), "date": day(r.get("coachingdate")),
                         "markdown": md, "owner": clean(r.get("supervisor")),
                         "action": clean(r.get("agreedaction")),
                         "follow_up": day(r["followupdate"]) if r.get("followupdate") else ""}
                coaching.append(entry)
            except (DataError, ValueError):
                issue(r, "Invalid coaching row; excluded from coaching comparisons")

    normalized = []
    for a in sorted(audits.values(), key=lambda a: (a["date"], a["audit_id"])):
        person = roster.get(a["agent_id"])
        if not person:
            person = {"agent": a['source_agent'] or sample_names.get(a['agent_id']) or a['agent_id'],
                      "supervisor": "Unassigned", "manager": "Unassigned", "site": "Unassigned"}
            issues.append({"file": a["source"], "row": a["source_row"], "issue": "Agent login is absent from the roster"})
        else:
            person = {**person, "agent": person['agent'] or a['source_agent'] or sample_names.get(a['agent_id']) or a['agent_id']}
            roster[a['agent_id']] = person
        a.update(person)
        a["dispositions"] = tracker.get((a["transaction_id"], a["agent_id"]), [])
        normalized.append(a)
    mode = "Synthetic demo" if synthetic else "Imported data"
    return {"audits": normalized, "roster": roster, "coaching": coaching,
            "selected_calls": calls, "issues": issues, "mode": mode,
            "demo_names_from_sample": source_names_applied, "demo_names_version": 1,
            "input_files": [p.name for p in sorted(input_dir.glob("*.xlsx"))],
            "source_dir": str(input_dir.resolve())}


def mean(rows):
    scores = [a["score"] for a in rows if a["score"] is not None]
    return sum(scores) / len(scores) if scores else None


def period(rows, start, end):
    return [a for a in rows if start <= a["date"] <= end]


def markdown_stats(rows, markdown):
    evaluated = [a for a in rows if a["markdowns"].get(markdown) is not None]
    misses = [a for a in evaluated if a["markdowns"][markdown]]
    points_confirmed = bool(evaluated) and all(a["deductions"].get(markdown) is not None for a in misses)
    total = sum(a["deductions"][markdown] for a in misses) if points_confirmed else None
    return {"markdown": markdown, **CATALOG[markdown], "evaluated": len(evaluated), "misses": len(misses),
            "rate": len(misses) / len(evaluated) * 100 if evaluated else None,
            "points_lost": total, "points_per_audit": total / len(evaluated) if total is not None else None,
            "audit_ids": [a["audit_id"] for a in misses]}


def windows(report_date, coaching_date=None):
    end = date.fromisoformat(report_date)
    if coaching_date:
        coaching_day = date.fromisoformat(coaching_date)
        length = min(14, (end - coaching_day).days)
        if length <= 0:
            return None
        recent_start, prior_end = end - timedelta(days=length - 1), coaching_day - timedelta(days=1)
        prior_start = coaching_day - timedelta(days=length)
    else:
        recent_start, prior_start, prior_end = (end - timedelta(days=n) for n in (13, 27, 14))
    return {"recent_start": recent_start.isoformat(), "recent_end": end.isoformat(),
            "prior_start": prior_start.isoformat(), "prior_end": prior_end.isoformat(),
            "basis": "Since coaching (up to 14 days)" if coaching_date else "Last two weeks vs prior two weeks"}


def status_for(rows, markdown, report_date, coaching=(), rules=None):
    rules = {**RULES, **(rules or {})}
    sessions = [c for c in coaching if c["markdown"] == markdown and c["date"] <= report_date]
    coaching_date = max((c["date"] for c in sessions), default=None)
    bounds = windows(report_date, coaching_date)
    if bounds is None:
        return {"markdown": markdown, "status": "Awaiting post-coaching data", "windows": None,
                "recent": None, "prior": None, "change_pp": None}
    recent = markdown_stats(period(rows, bounds["recent_start"], bounds["recent_end"]), markdown)
    prior = markdown_stats(period(rows, bounds["prior_start"], bounds["prior_end"]), markdown)
    change = recent["rate"] - prior["rate"] if recent["rate"] is not None and prior["rate"] is not None else None
    qa = mean(period(rows, report_date[:7] + "-01", report_date))
    if min(recent["evaluated"], prior["evaluated"]) < rules["min_audits"]:
        status = "No data" if not recent["evaluated"] and not prior["evaluated"] else "Low sample"
    elif ((qa is not None and qa < rules["qa_attention"]) or change >= rules["change_pp"]
          or (coaching_date and recent["rate"] >= rules["high_rate"] and change > -rules["change_pp"])):
        status = "Needs attention"
    elif change <= -rules["change_pp"]:
        status = "Improving"
    else:
        status = "Stable"
    return {"markdown": markdown, "status": status, "windows": bounds, "recent": recent,
            "prior": prior, "change_pp": change}


def focus_for(stats):
    candidates = [s for s in stats if s["misses"]]
    if not candidates:
        return None
    known = all(s["points_lost"] is not None for s in candidates)
    field = "points_lost" if known else "rate"
    ordered = sorted(candidates, key=lambda s: (-s[field], -s["impact"], s["markdown"]))
    return {**ordered[0], "basis": "Confirmed points lost" if known else "Markdown frequency"}


def grow(focus):
    if not focus:
        return {"goal": "Maintain the behaviors meeting the QA standard.",
                "reality": "No recorded markdown in this period; review sample size.",
                "options": "Review a successful audited call together.",
                "will": "Agree when to check the next completed audits."}
    md = focus["markdown"]
    actions = {"OIR": "Confirm the request, take ownership, and check that the resolution is understood.",
               "Silence": "Explain what you are checking and keep the customer informed during pauses.",
               "Mailing Address": "Use the approved address-confirmation steps and document the result.",
               "TPS": "Practice the approved three-point script verbiage.",
               "Notes": "Use the approved note format to capture the reason, action, and outcome."}
    return {"goal": f"Meet the approved {md} standard on reviewed calls.",
            "reality": f"{md} was flagged on {focus['misses']} of {focus['evaluated']} evaluated audits.",
            "options": actions[md],
            "will": "Agree a practical change and follow-up date; review the next completed audits."}


def build_report(data, report_date=None, filters=None, start_date=None, rules=None):
    rules = {**RULES, **(rules or {})}
    all_audits = data["audits"]
    report_date = report_date or max((a["date"] for a in all_audits), default=(date.today() - timedelta(days=1)).isoformat())
    report_date = day(report_date)
    end = date.fromisoformat(report_date)
    filters = {k: v for k, v in (filters or {}).items() if v and v != "All"}
    base = [a for a in all_audits if all(a.get(k) == v for k, v in filters.items()
                                        if k in ("supervisor", "manager", "call_type", "agent_id", "site"))]
    month_start = report_date[:7] + "-01"
    chosen_start = day(start_date) if start_date else month_start
    if chosen_start > report_date:
        raise DataError("The start date must be on or before the report date")
    current = period(base, month_start, report_date)
    visible = period(base, chosen_start, report_date)
    prev_month = end.replace(day=1) - timedelta(days=1)
    prev_end = prev_month.replace(day=min(end.day, calendar.monthrange(prev_month.year, prev_month.month)[1])).isoformat()
    previous = period(base, prev_end[:7] + "-01", prev_end)
    qa, prev_qa = mean(current), mean(previous)
    md_filter = filters.get("markdown")
    names = [md_filter] if md_filter in CATALOG else list(CATALOG)
    md_stats = [markdown_stats(visible, md) for md in names]
    focus = focus_for(md_stats)
    lookup = {a["audit_id"]: a for a in visible}
    agents = []
    logins = {a["agent_id"] for a in base}
    # Include roster agents with no audits unless a call-type selection defines an audited-only scope.
    if "call_type" not in filters:
        logins |= {login for login, p in data["roster"].items()
                   if all((login if k == "agent_id" else p.get(k)) == v for k, v in filters.items()
                          if k in ("supervisor", "manager", "site", "agent_id"))}
    for login in sorted(logins):
        history = [a for a in base if a["agent_id"] == login]
        month_rows = [a for a in current if a["agent_id"] == login]
        person = data["roster"].get(login, next((a for a in history), {}))
        agent_qa = mean(month_rows)
        scored = sum(a["score"] is not None for a in month_rows)
        alert = "Needs attention" if agent_qa is not None and agent_qa < rules["qa_attention"] else "On track" if agent_qa is not None else "No data"
        if scored and scored < rules["min_audits"]:
            alert += " / Low sample"
        sessions = [c for c in data["coaching"] if c["agent_id"] == login]
        statuses = [status_for(history, md, report_date, sessions, rules) for md in names]
        selected = period(history, chosen_start, report_date)
        agent_focus = focus_for([markdown_stats(selected, md) for md in names])
        agents.append({"agent_id": login, "agent": person.get("agent") or login,
                       "supervisor": person.get("supervisor", "Unassigned"), "manager": person.get("manager", "Unassigned"),
                       "qa": agent_qa, "audits": len(month_rows), "scored_audits": scored, "qa_alert": alert,
                       "statuses": statuses, "focus": agent_focus, "grow": grow(agent_focus),
                       "audit_ids": [a["audit_id"] for a in selected]})

    rollups = {}
    for field in ("supervisor", "manager", "site"):
        buckets = defaultdict(list)
        for a in current:
            buckets[a[field]].append(a)
        rollups[field] = [{"group": group, "qa": mean(items), "audits": len(items),
                           "scored_audits": sum(a["score"] is not None for a in items),
                           "audit_ids": [a["audit_id"] for a in items]}
                          for group, items in sorted(buckets.items())]
    heatmap = [{**markdown_stats([a for a in visible if a["call_type"] == call_type], md), "call_type": call_type}
               for call_type in sorted({a["call_type"] for a in visible}) for md in names]
    audited_txns = {(a["transaction_id"], a["agent_id"]) for a in all_audits if a["date"] <= report_date}
    sample = [{**s, "review_status": "Reviewed" if (s["transaction_id"], s["agent_id"]) in audited_txns else "Awaiting audit"}
              for s in data["selected_calls"]]
    detail = visible if not md_filter else [a for a in visible if a["markdowns"].get(md_filter)]
    weekly = period(base, (end - timedelta(days=6)).isoformat(), report_date)
    weekly_stats = [markdown_stats(weekly, md) for md in names]
    weekly_focus = focus_for(weekly_stats)
    summary = f"QA MTD is {qa:.1f}% across {len(current)} audits." if qa is not None else "No scored audits are available for QA MTD."
    if weekly_focus:
        summary += f" This week's focus is {weekly_focus['markdown']}, flagged on {weekly_focus['misses']} of {weekly_focus['evaluated']} evaluated audits."
    attention = sum('Needs attention' in a['qa_alert'] for a in agents)
    summary += f" {attention} {'agent has' if attention == 1 else 'agents have'} a QA attention alert."
    return {"report_date": report_date, "start_date": chosen_start, "month_start": month_start,
            "previous_start": prev_end[:7] + "-01", "previous_end": prev_end, "filters": filters,
            "mode": data["mode"], "qa": qa, "previous_qa": prev_qa,
            "demo_names_from_sample": data.get('demo_names_from_sample', False),
            "change_pp": qa - prev_qa if qa is not None and prev_qa is not None else None,
            "audit_count": len(current), "scored_audits": sum(a["score"] is not None for a in current),
            "previous_count": len(previous), "visible_count": len(visible), "agents": agents,
            "rollups": rollups, "markdowns": md_stats, "heatmap": heatmap,
            "focus": focus, "summary": summary, "weekly_start": (end - timedelta(days=6)).isoformat(),
            "sample": sample, "issues": data["issues"], "detail": detail, "rules": rules,
            "pivot_qa": current, "pivot_markdowns": visible,
            "coverage_start": min((a["date"] for a in base), default=None),
            "coverage_end": max((a["date"] for a in base), default=None),
            "snapshot_id": digest({"data": data, "date": report_date, "start": chosen_start,
                                   "filters": filters, "rules": rules}),
            "built_at": datetime.now(ZoneInfo(TIMEZONE)).isoformat(timespec="seconds")}


def audit_display(a):
    return {"Audit ID": a["audit_id"], "Transaction ID": a["transaction_id"], "Audit date": a["date"],
            "Agent": a["agent"], "Login ID": a["agent_id"], "Supervisor": a["supervisor"],
            "Manager": a["manager"], "Call type": a["call_type"], "QA score": a["score"],
            "Markdowns": ", ".join(m for m in CATALOG if a["markdowns"].get(m)),
            "Unknown results": ", ".join(m for m in CATALOG if a["markdowns"].get(m) is None),
            "Evaluator": a["evaluator"], "Comment": a["comment"],
            "Dispositions": "; ".join(f"{d['markdown']}: {d['l1']} / {d['l2']}" for d in a["dispositions"])}
