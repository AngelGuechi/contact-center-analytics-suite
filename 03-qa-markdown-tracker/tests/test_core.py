from __future__ import annotations

import copy
import importlib.util
import json
import sys
import tempfile
import unittest
from datetime import date, timedelta
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from qa_shared.core import CATALOG, DataError, build_report, load_sources, markdown_stats, status_for, windows
from qa_shared.demo import ensure_demo, ensure_sample, write_table
from qa_shared.excel_report import write_report
from openpyxl import load_workbook


def audit(i, date_value, score=90, agent="a", md=None, deductions=None):
    return {"audit_id": str(i), "transaction_id": "T" + str(i), "agent_id": agent, "agent": agent,
            "supervisor": "Team A", "manager": "Manager A", "site": "Site A", "date": date_value,
            "score": score, "markdowns": {m: m == md for m in CATALOG},
            "deductions": {m: (deductions or {}).get(m) for m in CATALOG}, "comment": "", "evaluator": "Evaluator",
            "dispositions": [], "call_type": "Credit", "revision": 1, "source": "test", "source_row": 2}


def dataset(audits, coaching=None):
    return {"audits": audits, "roster": {}, "coaching": coaching or [], "selected_calls": [],
            "issues": [], "mode": "Synthetic demo"}


class CalculationTests(unittest.TestCase):
    def test_improvement_uses_evaluated_denominators(self):
        rows = [audit(i, (date(2026, 8, 4) + timedelta(days=i)).isoformat(), md="OIR" if i < 4 else None) for i in range(10)]
        rows += [audit(i + 10, (date(2026, 8, 18) + timedelta(days=i)).isoformat(), md="OIR" if i == 0 else None) for i in range(10)]
        status = status_for(rows, "OIR", "2026-08-31")
        self.assertEqual(status["status"], "Improving")
        self.assertEqual(status["change_pp"], -30)
        self.assertEqual(status["prior"]["rate"], 40)
        self.assertEqual(status["recent"]["rate"], 10)

    def test_low_sample_and_qa_alert_are_separate(self):
        rows = [audit(1, "2026-08-20", score=70, md="OIR")]
        report = build_report(dataset(rows), "2026-08-31")
        self.assertEqual(report["agents"][0]["qa_alert"], "Needs attention / Low sample")
        self.assertEqual(report["agents"][0]["statuses"][0]["status"], "Low sample")

    def test_unknown_flag_not_counted_as_pass(self):
        rows = [audit(1, "2026-08-20", md="OIR"), audit(2, "2026-08-20")]
        rows[1]["markdowns"]["OIR"] = None
        stats = markdown_stats(rows, "OIR")
        self.assertEqual((stats["evaluated"], stats["rate"]), (1, 100))

    def test_no_data_differs_from_true_zero_score(self):
        self.assertIsNone(build_report(dataset([]), "2026-08-31")["qa"])
        self.assertEqual(build_report(dataset([audit(1, "2026-08-31", 0)]), "2026-08-31")["qa"], 0)

    def test_rollup_pools_audits(self):
        rows = [audit(1, "2026-08-20", 0)] + [audit(i + 2, "2026-08-20", 100, "b") for i in range(9)]
        report = build_report(dataset(rows), "2026-08-31")
        self.assertEqual(report["qa"], 90)
        self.assertEqual(report["rollups"]["supervisor"][0]["qa"], 90)

    def test_february_last_month_boundary(self):
        report = build_report(dataset([audit(1, "2026-02-28", 80), audit(2, "2026-03-31", 90)]), "2026-03-31")
        self.assertEqual(report["previous_end"], "2026-02-28")
        self.assertEqual(report["change_pp"], 10)

    def test_markdown_filter_does_not_bias_qa(self):
        rows = [audit(1, "2026-08-20", 60, md="OIR"), audit(2, "2026-08-20", 100)]
        report = build_report(dataset(rows), "2026-08-31", {"markdown": "OIR"})
        self.assertEqual(report["qa"], 80)
        self.assertEqual(len(report["detail"]), 1)

    def test_report_period_changes_markdowns_but_not_mtd(self):
        rows = [audit(1, "2026-08-02", 60, md="OIR"), audit(2, "2026-08-20", 100)]
        report = build_report(dataset(rows), "2026-08-31", start_date="2026-08-18")
        self.assertEqual(report["qa"], 80)
        self.assertEqual(report["markdowns"][0]["rate"], 0)

    def test_coaching_caps_both_windows_and_excludes_session_day(self):
        result = windows("2026-08-31", "2026-08-01")
        self.assertEqual(result["recent_start"], "2026-08-18")
        self.assertEqual(result["prior_start"], "2026-07-18")
        self.assertEqual(result["prior_end"], "2026-07-31")
        recent = windows("2026-08-31", "2026-08-28")
        self.assertEqual(recent["recent_start"], "2026-08-29")
        self.assertEqual(recent["prior_start"], "2026-08-25")
        self.assertIsNone(windows("2026-08-31", "2026-08-31"))

    def test_actual_deductions_and_unknown_impact(self):
        a = audit(1, "2026-08-20", md="OIR", deductions={"OIR": 35})
        self.assertEqual(markdown_stats([a], "OIR")["points_lost"], 35)
        a["deductions"]["OIR"] = None
        self.assertIsNone(markdown_stats([a], "OIR")["points_lost"])

    def test_missing_scores_do_not_poison_average(self):
        report = build_report(dataset([audit(1, "2026-08-20", None), audit(2, "2026-08-20", 80)]), "2026-08-31")
        self.assertEqual((report["qa"], report["audit_count"], report["scored_audits"]), (80, 2, 1))

    def test_sample_has_no_invented_qa(self):
        with tempfile.TemporaryDirectory() as temp:
            sample = ensure_sample(Path(temp)/'Synthetic_sample.ods')
            data = load_sources(temp, sample)
        report = build_report(data, "2026-08-31")
        self.assertEqual((len(report["sample"]), len({s["agent_id"] for s in report["sample"]})), (1512, 42))
        self.assertEqual(report["audit_count"], 0)
        self.assertTrue(all(s["review_status"] == "Awaiting audit" for s in report["sample"]))


class ImportTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.folder = Path(self.temp.name)
        self.headers = ["Audit ID", "Revision", "Audit Date", "Transaction ID", "Login ID", "Final QA Score", "Call Type", *CATALOG]

    def tearDown(self):
        self.temp.cleanup()

    def put(self, filename="Audits.xlsx", rows=None):
        write_table(self.folder / filename, self.headers,
                    rows or [["A1", 1, "2026-08-20", "T1", "a", 85, "Credit", "Yes", "No", "No", "No", "No"]])

    def test_duplicate_exports_and_tracker_do_not_multiply(self):
        self.put()
        self.put("Audits-copy.xlsx")
        write_table(self.folder / "SNOM3.xlsx", ["Transaction ID", "User ID", "Markdown", "L1", "L2"],
                    [["T1", "a", "OIR", "Resolution", "A"], ["T1", "a", "OIR", "Resolution", "A"],
                     ["T1", "a", "Notes", "Documentation", "B"]])
        data = load_sources(self.folder)
        self.assertEqual(len(data["audits"]), 1)
        self.assertEqual(len(data["audits"][0]["dispositions"]), 2)
        self.assertEqual(build_report(data, "2026-08-31")["audit_count"], 1)

    def test_revisions_replace_but_distinct_audits_remain(self):
        self.put()
        self.put("Audits-new.xlsx", [["A1", 2, "2026-08-20", "T1", "a", 95, "Credit", "No", "No", "No", "No", "No"],
                                       ["A2", 1, "2026-08-20", "T1", "a", 75, "Credit", "Yes", "No", "No", "No", "No"]])
        data = load_sources(self.folder)
        self.assertEqual(len(data["audits"]), 2)
        self.assertEqual(next(a["score"] for a in data["audits"] if a["audit_id"] == "A1"), 95)

    def test_conflicting_same_revision_stops_update(self):
        self.put()
        self.put("Audits-conflict.xlsx", [["A1", 1, "2026-08-20", "T1", "a", 95, "Credit", "No", "No", "No", "No", "No"]])
        with self.assertRaises(DataError):
            load_sources(self.folder)

    def test_excel_percentage_score_and_invalid_score(self):
        self.put(rows=[["A1", 1, "2026-08-20", "T1", "a", .85, "Credit", "No", "No", "No", "No", "No"],
                       ["A2", 1, "2026-08-20", "T2", "a", 170, "Credit", "No", "No", "No", "No", "No"]])
        wb = load_workbook(self.folder / "Audits.xlsx")
        wb.active["F2"].number_format = "0%"
        wb.save(self.folder / "Audits.xlsx")
        data = load_sources(self.folder)
        self.assertEqual(data["audits"][0]["score"], 85)
        self.assertIsNone(data["audits"][1]["score"])

    def test_wrong_snom_agent_is_not_attached(self):
        self.put()
        write_table(self.folder / "SNOM3.xlsx", ["Transaction ID", "User ID", "Markdown"], [["T1", "wrong-agent", "OIR"]])
        data = load_sources(self.folder)
        self.assertEqual(data["audits"][0]["dispositions"], [])
        self.assertTrue(any("no matching audit" in i["issue"] for i in data["issues"]))

    def test_bad_update_preserves_workbook(self):
        spec = importlib.util.spec_from_file_location("qa_limited_cli", ROOT / "01-limited-tools/run.py")
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        target = self.folder / "review.xlsx"
        target.write_bytes(b"previous-good-report")
        self.put()
        self.put("Audits-conflict.xlsx", [["A1", 1, "2026-08-20", "T1", "a", 95, "Credit", "No", "No", "No", "No", "No"]])
        self.assertEqual(module.main(["--input", str(self.folder), "--output", str(target)]), 1)
        self.assertEqual(target.read_bytes(), b"previous-good-report")

    def test_text_does_not_become_excel_formula(self):
        a = audit(1, "2026-08-20")
        a["comment"] = '=HYPERLINK("https://example.com","malicious")'
        report = build_report(dataset([a]), "2026-08-31")
        target = self.folder / "safe.xlsx"
        write_report(report, target)
        wb = load_workbook(target)
        self.assertNotEqual(wb["Audit Detail"]["N9"].data_type, "f")
        self.assertEqual(wb.sheetnames, ["Overview", "Pivot Tables", "Markdowns", "Audit Detail", "Pivot Source"])
        self.assertIsNotNone(wb["Overview"]["M29"].hyperlink)
        wb.close()


if __name__ == "__main__":
    unittest.main()
