import sys
import tempfile
import unittest
import xml.etree.ElementTree as ET
from pathlib import Path
from zipfile import ZipFile
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
from qa_shared.core import build_report, load_sources, selected_calls, CATALOG
from qa_shared.demo import write_table, ensure_demo, ensure_sample
from qa_shared.excel_report import write_report
from test_core import audit, dataset


class PivotTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.folder = Path(self.temp.name)
        self.sample = ensure_sample(self.folder/'Synthetic_sample.ods')

    def tearDown(self):
        self.temp.cleanup()

    def export(self, data, **options):
        report = build_report(data, '2026-08-31', **options)
        path = self.folder/'report.xlsx'
        write_report(report,path)
        return report,load_workbook(path),path

    def test_supplied_names_preserved_exactly(self):
        path = self.sample
        calls = selected_calls(path)
        with ZipFile(path) as z:
            root = ET.fromstring(z.read('content.xml'))
        ns = {'t':'urn:oasis:names:tc:opendocument:xmlns:table:1.0','x':'urn:oasis:names:tc:opendocument:xmlns:text:1.0'}
        sheet = next(t for t in root.findall('.//t:table',ns) if t.get('{'+ns['t']+'}name') == 'Selected Calls')
        first = sheet.findall('t:table-row',ns)[1].findall('t:table-cell',ns)[1]
        name = ' '.join(''.join(p.itertext()) for p in first.findall('x:p',ns)).strip()
        self.assertEqual(calls[0]['agent'],name)
        self.assertFalse(any(s['agent'].startswith('Sample agent ') for s in calls))

    def test_missing_roster_retains_exported_name(self):
        write_table(self.folder/'Audits.xlsx',['Audit ID','Audit Date','Transaction ID','Login ID','Agent','Final QA Score','Call Type',*CATALOG],
            [['A1','2026-08-21','T1','agent-1','María López',80,'Credit','Yes','No','No','No','No']])
        data = load_sources(self.folder)
        self.assertEqual(data['audits'][0]['agent'],'María López')
        self.assertEqual(data['audits'][0]['supervisor'],'Unassigned')

    def test_simulated_qa_uses_all_supplied_names_without_changing_results(self):
        folder = ensure_demo(self.folder/'simulation')
        baseline = build_report(load_sources(folder), '2026-08-31')
        sample = self.sample
        data = load_sources(folder, sample)
        report = build_report(data, '2026-08-31')
        names = list(dict((s['agent_id'],s['agent']) for s in data['selected_calls']).values())
        self.assertEqual([a['agent'] for a in report['agents']], names)
        self.assertTrue(report['demo_names_from_sample'])
        self.assertEqual((report['qa'],report['audit_count']), (baseline['qa'],baseline['audit_count']))
        self.assertEqual([a['statuses'] for a in report['agents']], [a['statuses'] for a in baseline['agents']])
        self.assertTrue(all(s['review_status']=='Awaiting audit' for s in report['sample']))
        path = self.folder/'source-names.xlsx'
        write_report(report,path)
        wb = load_workbook(path)
        try:
            self.assertEqual(wb['Overview']['B29'].value,names[0])
            self.assertEqual(wb['Markdowns']['B30'].value,names[0])
            self.assertEqual(wb['Audit Detail']['E9'].value,names[0])
            self.assertEqual(wb['Pivot Tables']['B29'].value,names[0]+' (DEMO001)')
            self.assertIn('SIMULATED QA RESULTS',wb['Overview']['B4'].value)
            agent_cache=wb['Pivot Tables']._pivots[1].cache.cacheFields[1].sharedItems
            self.assertTrue(all('Demo agent ' not in item.v for item in agent_cache._fields))
        finally:wb.close()

    def test_imported_roster_names_are_not_replaced_by_simulation_mapping(self):
        write_table(self.folder/'Roster.xlsx',['Login ID','Agent','Supervisor','Manager','Site'],
                    [['DEMO001','Original supplied roster name','Team','Manager','Site']])
        data=load_sources(self.folder,self.sample)
        self.assertEqual(data['roster']['DEMO001']['agent'],'Original supplied roster name')
        self.assertFalse(data['demo_names_from_sample'])

    def test_native_pivots_and_pooled_qa_cache(self):
        rows = [audit(1,'2026-08-20',0)] + [audit(i+2,'2026-08-20',100,'b') for i in range(9)]
        report,wb,path = self.export(dataset(rows))
        try:
            pivots = wb['Pivot Tables']._pivots
            self.assertEqual([p.name for p in pivots],['QATeams','QAAgents','MarkdownFrequency'])
            self.assertEqual(pivots[0].cache.recordCount,10)
            self.assertEqual(pivots[0].dataFields[0].subtotal,'average')
            self.assertAlmostEqual(wb['Pivot Tables']['C14'].value,report['qa']/100)
            self.assertEqual(wb['Pivot Tables']['D14'].value,10)
            self.assertEqual(wb['Pivot Source'].sheet_state,'hidden')
            self.assertTrue(all(wb[s]['E1'].hyperlink.target == "#'Pivot Tables'!B2" for s in ['Overview','Markdowns','Audit Detail']))
            with ZipFile(path) as z:
                self.assertEqual(len([n for n in z.namelist() if n.startswith('xl/pivotTables/pivotTable') and n.endswith('.xml')]),3)
                self.assertGreater(len([n for n in z.namelist() if 'pivotCacheRecords' in n and n.endswith('.xml')]),0)
        finally:
            wb.close()

    def test_markdown_filter_keeps_qa_population_and_ignores_unknowns(self):
        rows = [audit(1,'2026-08-20',60,md='OIR'),audit(2,'2026-08-20',100),audit(3,'2026-08-20',None)]
        rows[2]['markdowns']['OIR'] = None
        report,wb,_ = self.export(dataset(rows),filters={'markdown':'OIR'})
        try:
            self.assertEqual(wb['Pivot Tables']._pivots[0].cache.recordCount,3)
            self.assertEqual(wb['Pivot Tables']._pivots[2].cache.recordCount,3)
            self.assertEqual(wb['Pivot Tables']['J17'].value,.5)
            self.assertEqual(wb['Pivot Tables']['K17'].value,1)
            self.assertEqual(wb['Pivot Tables']['L17'].value,2)
            self.assertEqual(wb['Pivot Tables']['C14'].value,.8)
        finally:
            wb.close()

    def test_sample_pivot_without_fake_audits(self):
        data = load_sources(self.folder,self.sample)
        report,wb,_ = self.export(data)
        try:
            pivots = wb['Pivot Tables']._pivots
            self.assertEqual([p.name for p in pivots],['SampleCoverage'])
            self.assertEqual(pivots[0].cache.recordCount,1512)
            self.assertEqual(len(pivots[0].rowItems)-1,42)
            self.assertIsNone(report['qa'])
        finally:
            wb.close()

    def test_same_names_with_different_ids_stay_distinct(self):
        rows=[audit(1,'2026-08-20',80,'a'),audit(2,'2026-08-20',90,'b')]
        for a in rows:a['agent']='Alex Smith'
        _,wb,_ = self.export(dataset(rows))
        try:
            self.assertEqual(len(wb['Pivot Tables']._pivots[1].rowItems)-1,2)
        finally:wb.close()


if __name__ == '__main__':
    unittest.main()
