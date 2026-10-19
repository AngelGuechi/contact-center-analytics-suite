"""File-backed integration checks through LocalStore's public operations."""
from __future__ import annotations

import hashlib
import json
import shutil
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from avoidance.store import LocalStore, safe_relative

FIXTURES = Path(__file__).resolve().parents[1] / "fixtures"


class LocalStoreIntegrationTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.seed_temp = tempfile.TemporaryDirectory()
        cls.seed_root = Path(cls.seed_temp.name) / "store"
        seed_store = LocalStore(cls.seed_root)
        seed_store.begin_job("Import fixture sources")
        cls.initial_report = seed_store.run_job(seed_store.import_folder, FIXTURES)

    @classmethod
    def tearDownClass(cls):
        cls.seed_temp.cleanup()

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name) / "store"
        self.input = Path(self.temp.name) / "input"
        shutil.copytree(self.seed_root, self.root)
        shutil.copytree(FIXTURES, self.input)
        self.store = LocalStore(self.root)

    def tearDown(self):
        self.temp.cleanup()

    def import_now(self, source=None):
        self.store.begin_job("Import integration test sources")
        return self.store.run_job(self.store.import_folder, source or self.input)

    def ana_slice(self, report=None):
        report = report or self.store.report()
        return next(row for row in report["daily"]["slices"]
            if row["agent_id"] == "123654" and row["call_type"] == "General Billing")

    def set_manifest(self, **changes):
        path = self.input / "manifest.json"
        manifest = json.loads(path.read_text(encoding="utf-8"))
        manifest.update(changes)
        path.write_text(json.dumps(manifest, indent=2), encoding="utf-8")

    def edit_tool2(self, changes, day="2026-10-06"):
        path = self.input / "Tool2" / day / f"tool2_{day}.xlsx"
        book = load_workbook(path)
        for address, value in changes.items():
            book.active[address] = value
        book.save(path)
        book.close()

    def test_initial_native_import_publishes_expected_patterns_and_pages(self):
        report = self.store.report()
        self.assertEqual(report["overview"]["daily_patterns"], 1)
        self.assertEqual(report["overview"]["intraday_leads"], 1)
        ana = self.ana_slice(report)
        self.assertEqual((ana["signals"], ana["transactions"]), (6, 30))
        self.assertEqual(ana["reasons"], ["peer_outlier", "self_change"])
        detail = self.store.evidence(ana["evidence_id"], page=2, page_size=7)
        self.assertEqual(detail["transaction_count"], 30)
        self.assertEqual(len(detail["transactions"]), 7)
        self.assertTrue(all(row["sources"][0]["raw_cells"] for row in detail["transactions"]))
        self.assertNotIn("evidence", report)
        json.dumps(report, allow_nan=False)

    def test_repeat_import_uses_normalized_cache_and_same_publication(self):
        before = self.store.report()["publication"]["id"]
        with patch("avoidance.store.read_dataset", side_effect=AssertionError("Reparsed cached input")), patch("avoidance.store.load_dataset", side_effect=AssertionError("Reloaded unchanged normalized data")), patch("avoidance.store.build_report", side_effect=AssertionError("Rebuilt unchanged report")):
            after = self.import_now()
        self.assertEqual(after["publication"]["id"], before)
        self.assertIn("reused", self.store.status()["job"]["message"].lower())
        attempts = [json.loads(line) for line in (self.root / "attempts.jsonl").read_text(encoding="utf-8").splitlines()]
        self.assertTrue(attempts[-1]["reused_normalized"])

    def test_default_import_restores_manifest_date_after_explicit_other_date(self):
        earlier = self.store.import_folder(self.input, "2026-10-05")
        self.assertEqual(earlier["report_date"], "2026-10-05")
        default = self.store.import_folder(self.input)
        self.assertEqual(default["report_date"], "2026-10-06")

    def test_retained_sources_survive_original_change_without_changing_report(self):
        before = self.store.report()["publication"]["id"]
        record = self.store.evidence(self.ana_slice()["evidence_id"])["transactions"][0]
        locator = record["sources"][0]
        retained = list((self.root / "sources").glob("*/" + locator["file"]))
        self.assertEqual(len(retained), 1)
        self.assertEqual(hashlib.sha256(retained[0].read_bytes()).hexdigest(), locator["sha256"])
        self.edit_tool2({"G2": 7})
        self.assertEqual(hashlib.sha256(retained[0].read_bytes()).hexdigest(), locator["sha256"])
        self.assertEqual(self.store.report()["publication"]["id"], before)

    def test_human_review_survives_store_restart_and_remains_append_only(self):
        evidence_id = self.ana_slice()["evidence_id"]
        first = self.store.save_review(evidence_id, "Technical issue", "Synthetic reviewer", "Inspect the reported synthetic incident.")
        second = self.store.save_review(evidence_id, "Needs review", "Synthetic reviewer", "Await source confirmation.")
        restarted = LocalStore(self.root)
        self.assertEqual([row["id"] for row in restarted.reviews(evidence_id)], [first["id"], second["id"]])
        self.assertEqual(restarted.report()["reviews"][evidence_id]["state"], "Needs review")
        self.assertEqual(first["transaction_ids"], restarted.evidence(evidence_id)["transaction_ids"])
        self.assertTrue(first["timestamp"])
        self.assertTrue(first["settings"])

    def test_new_rules_create_new_evidence_and_keep_old_review_as_context(self):
        old_slice = self.ana_slice()
        old_version = self.store.report()["settings"]["version"]
        decision = self.store.save_review(old_slice["evidence_id"], "Legitimate interaction", "Synthetic reviewer", "Checked original evidence.")
        self.store.begin_job("Change demo threshold")
        updated = self.store.run_job(self.store.update_settings, {"daily_min_rate": 0.15})
        new_slice = self.ana_slice(updated)
        self.assertNotEqual(new_slice["evidence_id"], old_slice["evidence_id"])
        self.assertNotEqual(updated["settings"]["version"], old_version)
        detail = self.store.evidence(new_slice["evidence_id"])
        self.assertEqual(detail["reviews"], [])
        self.assertIn(decision["id"], [row["id"] for row in detail["earlier_reviews"]])
        self.assertNotIn(new_slice["evidence_id"], updated["reviews"])
        restarted = LocalStore(self.root)
        self.assertEqual(restarted.report()["settings"], updated["settings"])

    def test_bad_replacement_keeps_last_good_publication_and_reviews(self):
        old_id = self.store.report()["publication"]["id"]
        evidence_id = self.ana_slice()["evidence_id"]
        decision = self.store.save_review(evidence_id, "Needs review", "Synthetic reviewer", "Pending source validation.")
        self.edit_tool2({"E3": -10})
        result = self.import_now()
        self.assertEqual(result["publication"]["id"], old_id)
        self.assertEqual(self.store.status()["job"]["state"], "failed")
        self.assertEqual(self.store.report()["reviews"][evidence_id]["id"], decision["id"])
        self.assertTrue(self.store.status()["attempts"][-1]["issues"])

    def test_missing_prior_day_publishes_insufficient_history_without_false_flag(self):
        path = self.input / "Tool2/2026-09-25/tool2_2026-09-25.xlsx"
        path.unlink()
        report = self.import_now()
        self.assertEqual(report["overview"]["daily_patterns"], 0)
        ana = self.ana_slice(report)
        self.assertEqual((ana["signals"], ana["transactions"]), (6, 30))
        self.assertEqual(ana["reasons"], [])
        self.assertIn("2026-09-25", ana["self_baseline"]["missing_days"])
        self.assertTrue(any(row["code"] == "missing_file" for row in report["import_issues"]))

    def test_corrupt_export_preserves_previous_report_and_visible_failure(self):
        before = self.store.report()["publication"]["id"]
        path = self.input / "Tool2/2026-10-06/tool2_2026-10-06.xlsx"
        path.write_bytes(b"not an Excel workbook")
        report = self.import_now()
        self.assertEqual(report["publication"]["id"], before)
        self.assertEqual(report["job"]["state"], "failed")

    def test_safe_unmapped_records_remain_visible_after_publication(self):
        self.edit_tool2({"I2": "UNKNOWN_LOGIN", "L3": 77777})
        report = self.import_now()
        dataset = self.store.dataset()
        self.assertTrue(any(row["agent_id"] == "unmapped:UNKNOWN_LOGIN" for row in dataset["transactions"]))
        self.assertTrue(any(row["call_type"] == "Unknown call type" for row in dataset["transactions"]))
        self.assertTrue({"unmapped_login", "unmapped_queue"}.issubset({row["code"] for row in report["import_issues"]}))

    def test_store_rejects_mixing_synthetic_and_operational_data(self):
        before = self.store.report()["publication"]["id"]
        self.set_manifest(synthetic=False)
        with self.assertRaises(ValueError):
            self.import_now()
        self.assertEqual(self.store.report()["publication"]["id"], before)
        self.assertEqual(self.store.status()["job"]["state"], "failed")

    def test_operational_store_cannot_later_load_synthetic_demo(self):
        operational_root = Path(self.temp.name) / "operational"
        store = LocalStore(operational_root)
        self.set_manifest(synthetic=False)
        store.import_folder(self.input)
        before = store.report()["publication"]["id"]
        with self.assertRaises(ValueError):
            store.import_folder(FIXTURES)
        self.assertEqual(store.report()["publication"]["id"], before)

    def test_unsupported_review_does_not_append_log(self):
        evidence_id = self.ana_slice()["evidence_id"]
        for state, reviewer, note in (("Automatic discipline", "Demo", "Disallowed state"),
            ("Needs review", "  ", "A note"), ("Needs review", "Demo", "  ")):
            with self.subTest(state=state, reviewer=reviewer, note=note):
                with self.assertRaises(ValueError):
                    self.store.save_review(evidence_id, state, reviewer, note)
        self.assertEqual(self.store.reviews(), [])

    def test_settings_reject_unsupported_and_nonfinite_values_atomically(self):
        before = self.store.report()["publication"]["id"]
        old_settings = dict(self.store.settings)
        invalid = [
            {"unknown_setting": 1}, {"version": "user-invented-version"},
            {"daily_min_rate": float("nan")}, {"daily_min_rate": float("inf")},
            {"short_handle_seconds": float("inf")},
            {"short_handle_seconds": {"General Billing": float("nan")}},
            {"short_handle_seconds": {"General Billing": True}},
            {"daily_min_rate": 1.01}, {"daily_min_rate": -0.01},
            {"rate_increase_pp": 101}, {"daily_min_transactions": 0},
            {"daily_min_transactions": 2.5}, {"daily_min_transactions": True},
            {"interval_min_minutes": 100, "interval_max_minutes": 90},
        ]
        for changes in invalid:
            with self.subTest(changes=changes):
                with self.assertRaises(ValueError):
                    self.store.update_settings(changes)
                self.assertEqual(self.store.settings, old_settings)
                self.assertEqual(self.store.report()["publication"]["id"], before)

    def test_another_job_cannot_start_while_import_is_running(self):
        self.store.begin_job("First import")
        with self.assertRaises(ValueError):
            self.store.begin_job("Second import")
        self.assertEqual(self.store.status()["job"]["state"], "running")

    def test_every_daily_slice_has_exact_evidence_after_restart(self):
        restarted = LocalStore(self.root)
        report = restarted.report()
        labels = set()
        for row in report["daily"]["slices"]:
            with self.subTest(agent=row["agent_id"]):
                detail = restarted.evidence(row["evidence_id"], publication_id=report["publication"]["id"])
                self.assertEqual(detail["transaction_count"], row["transactions"])
                self.assertEqual(detail["population_count"], row["transactions"])
                self.assertEqual(len(detail["transactions"]), row["transactions"])
                self.assertEqual([record["transaction_id"] for record in detail["transactions"]], row["transaction_ids"])
                self.assertTrue(all(record["sources"][0]["raw_cells"] for record in detail["transactions"]))
                self.assertTrue(detail["is_current"])
                labels.add(row["label"])
        self.assertIn("Low sample", labels)
        self.assertIn("No qualifying daily pattern", labels)
        empty = next(agent for agent in report["daily"]["agents"] if not agent["transactions"])
        self.assertEqual(empty["label"], "No transactions")
        self.assertIsNone(empty["rate"])
        self.assertEqual(empty["evidence_ids"], [])

    def test_exact_archived_publication_keeps_its_source_rows_after_new_import(self):
        old_report = self.store.report()
        old_row = self.ana_slice(old_report)
        old_publication = old_report["publication"]["id"]
        original = self.store.evidence(old_row["evidence_id"])["transactions"][0]
        decision = self.store.save_review(old_row["evidence_id"], "Needs review", "Reviewer", "Original source checked.")
        self.edit_tool2({"H2": original["holds"] + 1})
        new_report = self.import_now()
        new_row = self.ana_slice(new_report)
        self.assertNotEqual(old_row["evidence_id"], new_row["evidence_id"])
        with self.assertRaises(KeyError):
            self.store.evidence(old_row["evidence_id"])
        restarted = LocalStore(self.root)
        archived = restarted.evidence(old_row["evidence_id"], publication_id=old_publication)
        current = restarted.evidence(new_row["evidence_id"], publication_id=new_report["publication"]["id"])
        self.assertFalse(archived["is_current"])
        self.assertEqual(archived["publication_id"], old_publication)
        self.assertEqual(archived["current_publication_id"], new_report["publication"]["id"])
        self.assertEqual(archived["report_date"], old_report["report_date"])
        self.assertEqual(archived["transactions"][0], original)
        self.assertEqual(current["transactions"][0]["holds"], original["holds"] + 1)
        self.assertNotEqual(current["transactions"][0]["sources"][0]["sha256"], original["sources"][0]["sha256"])
        self.assertEqual(archived["reviews"][0]["id"], decision["id"])
        self.assertTrue(current["is_current"])
        reviews_before = restarted.reviews()
        with self.assertRaises(KeyError):
            restarted.save_review(old_row["evidence_id"], "Technical issue", "Reviewer", "Stale decision must not attach to the new report.")
        self.assertEqual(restarted.reviews(), reviews_before)

    def test_short_record_filter_precedes_pagination_for_every_population(self):
        row = self.ana_slice()
        current_ids = []
        for page in (1, 2, 3):
            detail = self.store.evidence(row["evidence_id"], page=page, page_size=2, signals_only=True)
            self.assertEqual((detail["transaction_count"], detail["population_count"]), (6, 30))
            self.assertTrue(detail["signals_only"])
            self.assertEqual(len(detail["transactions"]), 2)
            current_ids.extend(record["transaction_id"] for record in detail["transactions"])
        self.assertEqual(current_ids, row["signal_transaction_ids"])
        for scope, count, population in (("peer", 56, 2800), ("self", 14, 700)):
            detail = self.store.evidence(row["evidence_id"], page=2, page_size=5, scope=scope, signals_only=True)
            self.assertEqual((detail["transaction_count"], detail["population_count"]), (count, population))
            self.assertEqual(len(detail["transactions"]), 5)
            self.assertTrue(all(record["duration_seconds"] <= 15 and record["transfers"] == 0 for record in detail["transactions"]))
            self.assertEqual(detail["signal_transaction_ids"], [record["transaction_id"] for record in detail["transactions"]])
        with self.assertRaises(ValueError):
            self.store.evidence(self.store.report()["intraday"]["leads"][0]["evidence_id"], signals_only=True)

    def test_missing_cache_recovers_original_records_without_writing_runtime_files(self):
        cache = self.root / self.store._publication["normalized"]
        cache.unlink()
        before = {path.relative_to(self.root): hashlib.sha256(path.read_bytes()).hexdigest()
                  for path in self.root.rglob("*") if path.is_file()}
        restarted = LocalStore(self.root)
        row = self.ana_slice(restarted.report())
        detail = restarted.evidence(row["evidence_id"])
        self.assertEqual((detail["transaction_count"], len(detail["transactions"])), (30, 30))
        self.assertTrue(all(record["sources"][0]["raw_cells"] for record in detail["transactions"]))
        self.assertFalse(cache.exists())
        after = {path.relative_to(self.root): hashlib.sha256(path.read_bytes()).hexdigest()
                 for path in self.root.rglob("*") if path.is_file()}
        self.assertEqual(after, before)

    def test_corrupt_cache_recovers_in_memory_and_changed_retained_sources_are_rejected(self):
        cache = self.root / self.store._publication["normalized"]
        cache.write_text("{", encoding="utf-8")
        restarted = LocalStore(self.root)
        detail = restarted.evidence(self.ana_slice()["evidence_id"])
        self.assertEqual(len(detail["transactions"]), 30)
        self.assertEqual(cache.read_text(encoding="utf-8"), "{")
        archived_workbook = self.root / self.store._publication["sources"] / "Tool2/2026-10-06/tool2_2026-10-06.xlsx"
        book = load_workbook(archived_workbook)
        book.active["G2"] = 1
        book.save(archived_workbook)
        book.close()
        changed = LocalStore(self.root)
        with self.assertRaisesRegex(ValueError, "no longer match the saved evidence"):
            changed.evidence(self.ana_slice()["evidence_id"])
        self.assertEqual(changed.report()["publication"]["id"], self.store.report()["publication"]["id"])

    def test_saved_publication_reference_cannot_escape_store(self):
        evidence_id = self.ana_slice()["evidence_id"]
        for publication_id in ("../private", "C:\\private", "bad-id", "a" * 63):
            with self.subTest(publication_id=publication_id), self.assertRaises(ValueError):
                self.store.evidence(evidence_id, publication_id=publication_id)
        with self.assertRaises(KeyError):
            self.store.evidence(evidence_id, publication_id="0" * 64)

    def test_review_checks_publication_even_when_transaction_evidence_is_unchanged(self):
        before = self.store.report()
        evidence_id = self.ana_slice(before)["evidence_id"]
        old_publication = before["publication"]["id"]
        self.set_manifest(annotations=[{"annotation_id": "context-update", "note": "Updated site context for human review."}])
        after = self.import_now()
        self.assertNotEqual(after["publication"]["id"], old_publication)
        self.assertEqual(self.ana_slice(after)["evidence_id"], evidence_id)
        reviews_before = self.store.reviews()
        with self.assertRaisesRegex(ValueError, "source report changed"):
            self.store.save_review(evidence_id, "Needs review", "Reviewer", "Decision from a previous publication.", publication_id=old_publication)
        self.assertEqual(self.store.reviews(), reviews_before)
        accepted = self.store.save_review(evidence_id, "Needs review", "Reviewer", "Checked the current source context.", publication_id=after["publication"]["id"])
        self.assertEqual(accepted["publication_id"], after["publication"]["id"])


class RelativePathTests(unittest.TestCase):
    def test_relative_input_paths(self):
        self.assertEqual(safe_relative("Tool2/2026-10-06/file.xlsx"), Path("Tool2/2026-10-06/file.xlsx"))
        self.assertEqual(safe_relative("Tool2\\2026-10-06\\file.xlsx"), Path("Tool2/2026-10-06/file.xlsx"))
        for raw in ("../private.xlsx", "/private.xlsx", "C:\\private.xlsx", "a/../b", "a//b", "", 7, {}):
            with self.subTest(raw=raw):
                self.assertRaises(ValueError, safe_relative, raw)


if __name__ == "__main__":
    unittest.main()
