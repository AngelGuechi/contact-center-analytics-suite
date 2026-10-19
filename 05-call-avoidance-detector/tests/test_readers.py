"""Native format, atomic import, mapping, completeness and lineage checks."""
from __future__ import annotations

import json
import shutil
import tempfile
import unittest
from datetime import datetime, time, timedelta
from pathlib import Path

from openpyxl import load_workbook

from avoidance.fixtures import generate_demo
from avoidance.readers import duration_seconds, normalize_agent_id, read_dataset, rebuild_timestamp


class NativeReaderTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.base_temp = tempfile.TemporaryDirectory()
        cls.base = generate_demo(Path(cls.base_temp.name) / "source")

    @classmethod
    def tearDownClass(cls):
        cls.base_temp.cleanup()

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.input = Path(self.temp.name) / "inputs"
        shutil.copytree(self.base, self.input)

    def tearDown(self):
        self.temp.cleanup()

    def manifest(self):
        return json.loads((self.input / "manifest.json").read_text(encoding="utf-8"))

    def tool2(self, day="2026-10-06"):
        return self.input / "Tool2" / day / f"tool2_{day}.xlsx"

    def edit(self, path, changes, sheet_name=None):
        book = load_workbook(path)
        sheet = book[sheet_name] if sheet_name else book.active
        for address, value in changes.items():
            sheet[address] = value
        book.save(path)
        book.close()

    def test_demo_import_and_planted_source_numbers(self):
        dataset = read_dataset(self.input)
        self.assertEqual(dataset["issues"], [])
        self.assertEqual(len(dataset["transactions"]), 3630)
        self.assertEqual(len(dataset["complete_days"]), 15)
        ana = [row for row in dataset["transactions"] if row["agent_id"] == "123654"]
        current = [row for row in ana if row["business_date"] == "2026-10-06"]
        past = [row for row in ana if row["business_date"] != "2026-10-06"]
        self.assertEqual(len(current), 30)
        self.assertEqual(sum(row["duration_seconds"] <= 15 and row["transfers"] == 0 for row in current), 6)
        self.assertEqual(len(past), 700)
        self.assertEqual(sum(row["duration_seconds"] <= 15 and row["transfers"] == 0 for row in past), 14)
        bruno = [(snap["captured_at"], next(row for row in snap["agents"] if row["agent_id"] == "235694"))
            for snap in dataset["snapshots"][:3]]
        self.assertEqual([row["calls"] for _, row in bruno], [10, 15, 20])
        self.assertEqual([row["talk_seconds"] for _, row in bruno], [3000, 3120, 3270])
        self.assertEqual(bruno[0][0], "2026-10-06T15:00:00Z")
        self.assertEqual(dataset["import_summary"]["blank_rows"], 15)
        self.assertEqual(dataset["import_summary"]["skipped_rows"], 4)

    def test_42_column_tool2_and_side_by_side_roster(self):
        book = load_workbook(self.tool2(), read_only=True)
        self.assertEqual(book.active.max_column, 42)
        book.close()
        dataset = read_dataset(self.input)
        self.assertEqual({row["agent_id"] for row in dataset["roster"]},
            {"123654", "235694", "345678", "456789", "012345", "567890"})
        self.assertTrue(all(row["call_type"] == "General Billing" for row in dataset["transactions"]))

    def test_source_lineage_keeps_native_cells_and_hash(self):
        record = next(row for row in read_dataset(self.input)["transactions"] if row["business_date"] == "2026-10-06")
        source = record["sources"][0]
        self.assertEqual(source["row"], 2)
        self.assertEqual(source["sheet"], "Export")
        self.assertEqual(len(source["sha256"]), 64)
        self.assertEqual(source["import_id"], "xlsx:" + source["sha256"])
        self.assertIn("Time(Server)", source["raw_cells"])
        self.assertEqual(len(source["raw_cells"]), 42)
        self.assertEqual(source["raw_cells"]["Login ID"]["value"], "SYNANA")
        json.dumps(source)  # persisted dataset has no non-JSON Excel objects

    def test_exact_duplicate_merges_lineage_without_double_counting(self):
        duplicate = self.input / "zz_duplicate.xlsx"
        shutil.copy2(self.tool2(), duplicate)
        dataset = read_dataset(self.input)
        self.assertEqual(len(dataset["transactions"]), 3630)
        self.assertEqual(dataset["import_summary"]["duplicate_transactions"], 130)
        current = [row for row in dataset["transactions"] if row["business_date"] == "2026-10-06"]
        self.assertTrue(all(len(row["sources"]) == 2 for row in current))
        self.assertIn("2026-10-06", dataset["complete_days"])

    def test_conflicting_duplicate_rejects_entire_new_file(self):
        duplicate = self.input / "zz_conflicting.xlsx"
        shutil.copy2(self.tool2(), duplicate)
        self.edit(duplicate, {"G2": 4})
        dataset = read_dataset(self.input)
        self.assertEqual(len(dataset["transactions"]), 3630)
        self.assertEqual(dataset["import_summary"]["files_rejected"], 1)
        self.assertNotIn("2026-10-06", dataset["complete_days"])
        self.assertIn("2026-10-05", dataset["complete_days"])
        self.assertTrue(any("Conflicting duplicate" in item["message"] for item in dataset["issues"]))

    def test_missing_required_day_marks_history_incomplete(self):
        self.tool2("2026-09-25").unlink()
        dataset = read_dataset(self.input)
        self.assertEqual(dataset["import_summary"]["files_missing"], 1)
        self.assertNotIn("2026-09-25", dataset["complete_days"])
        self.assertIn("2026-09-25", dataset["expected_days"])
        self.assertEqual(len(dataset["complete_days"]), 14)

    def test_filename_and_rows_alone_do_not_establish_completeness(self):
        (self.input / "manifest.json").unlink()
        dataset = read_dataset(self.input)
        self.assertEqual(len(dataset["transactions"]), 3630)
        self.assertEqual(dataset["complete_days"], [])
        self.assertTrue(any(item["code"] == "manifest_unavailable" for item in dataset["issues"]))

    def test_completeness_needs_matching_accepted_declared_file(self):
        manifest = self.manifest()
        manifest["files"] = [item for item in manifest["files"] if item.get("business_date") != "2026-09-24"]
        (self.input / "manifest.json").write_text(json.dumps(manifest), encoding="utf-8")
        dataset = read_dataset(self.input)
        self.assertNotIn("2026-09-24", dataset["complete_days"])
        self.assertEqual(len(dataset["transactions"]), 3630)

    def test_invalid_row_rejects_workbook_preserving_other_days(self):
        self.edit(self.tool2(), {"E3": -5})
        dataset = read_dataset(self.input)
        self.assertEqual(len(dataset["transactions"]), 3500)
        self.assertNotIn("2026-10-06", dataset["complete_days"])
        self.assertEqual(len(dataset["complete_days"]), 14)
        self.assertTrue(any(item["row"] == 3 and item["code"] == "invalid_transaction" for item in dataset["issues"]))

    def test_unlabelled_invalid_workbook_uses_own_row_day_for_incompleteness(self):
        duplicate = self.input / "unlabelled.xlsx"
        shutil.copy2(self.tool2(), duplicate)
        self.edit(duplicate, {"E3": -5})
        dataset = read_dataset(self.input)
        self.assertEqual(len(dataset["transactions"]), 3630)
        self.assertNotIn("2026-10-06", dataset["complete_days"])
        self.assertEqual(len(dataset["complete_days"]), 14)

    def test_end_duration_and_minute_clock_are_validated(self):
        self.edit(self.tool2(), {"C2": "59.30.0"})
        dataset = read_dataset(self.input)
        self.assertEqual(dataset["import_summary"]["files_rejected"], 1)
        self.assertTrue(any("does not match" in item["message"] for item in dataset["issues"]))

    def test_unknown_login_and_queue_are_safe_and_visible(self):
        self.edit(self.tool2(), {"I2": "UNKNOWN_LOGIN", "L3": 77777})
        dataset = read_dataset(self.input)
        self.assertEqual(len(dataset["transactions"]), 3630)
        self.assertIn("2026-10-06", dataset["complete_days"])
        unknown = next(row for row in dataset["roster"] if row["agent_id"] == "unmapped:UNKNOWN_LOGIN")
        self.assertEqual(unknown["supervisor"], "Unassigned")
        self.assertEqual(unknown["name"], "Unassigned")
        self.assertEqual(sum(row["call_type"] == "Unknown call type" for row in dataset["transactions"]), 1)
        self.assertTrue({"unmapped_login", "unmapped_queue"}.issubset({item["code"] for item in dataset["issues"]}))

    def test_bad_snapshot_retains_rejected_time_to_break_streak(self):
        path = self.input / "Tool1/2026-10-06/tool1_2026-10-06_0930.xlsx"
        self.edit(path, {"X3": timedelta(seconds=999)}, "LiveAgentData")
        dataset = read_dataset(self.input)
        rejected = [snap for snap in dataset["snapshots"] if not snap["accepted"]]
        self.assertEqual(len(rejected), 1)
        self.assertEqual(rejected[0]["captured_at"], "2026-10-06T15:30:00Z")
        self.assertEqual(rejected[0]["agents"], [])
        self.assertTrue(rejected[0]["issues"])

    def test_missing_snapshot_retains_gap_time(self):
        path = self.input / "Tool1/2026-10-06/tool1_2026-10-06_0930.xlsx"
        path.unlink()
        dataset = read_dataset(self.input)
        gap = next(snap for snap in dataset["snapshots"] if not snap["accepted"])
        self.assertEqual(gap["captured_at"], "2026-10-06T15:30:00Z")
        self.assertEqual(gap["issues"][0]["code"], "missing_snapshot")

    def test_duplicate_snapshot_merges_rows_without_extra_interval(self):
        original = self.input / "Tool1/2026-10-06/tool1_2026-10-06_0930.xlsx"
        shutil.copy2(original, self.input / "zz_tool1_2026-10-06_0930.xlsx")
        dataset = read_dataset(self.input)
        self.assertEqual(len(dataset["snapshots"]), 4)
        self.assertEqual(dataset["import_summary"]["duplicate_snapshots"], 1)
        copied = next(snap for snap in dataset["snapshots"] if snap["captured_at"] == "2026-10-06T15:30:00Z")
        self.assertTrue(all(len(row["sources"]) == 2 for row in copied["agents"]))

    def test_conflicting_snapshot_keeps_invalid_time_in_chain(self):
        original = self.input / "Tool1/2026-10-06/tool1_2026-10-06_0930.xlsx"
        duplicate = self.input / "zz_tool1_2026-10-06_0930.xlsx"
        shutil.copy2(original, duplicate)
        self.edit(duplicate, {"T3": 4}, "LiveAgentData")
        dataset = read_dataset(self.input)
        copied = [snap for snap in dataset["snapshots"] if snap["captured_at"] == "2026-10-06T15:30:00Z"]
        self.assertTrue(copied)
        self.assertTrue(all(not snap["accepted"] for snap in copied))

    def test_declared_complete_day_also_needs_expected_day(self):
        manifest = self.manifest()
        manifest["expected_days"].remove("2026-10-06")
        (self.input / "manifest.json").write_text(json.dumps(manifest), encoding="utf-8")
        self.assertNotIn("2026-10-06", read_dataset(self.input)["complete_days"])

    def test_manifest_rejects_relative_or_absolute_escape(self):
        outside = Path(self.temp.name) / "outside.xlsx"
        shutil.copy2(self.tool2(), outside)
        for filename in ("../outside.xlsx", str(outside.resolve())):
            with self.subTest(filename=filename):
                manifest = self.manifest()
                manifest["files"].append(dict(file=filename, kind="tool2", required=True,
                    business_date="2026-10-06", complete=True))
                (self.input / "manifest.json").write_text(json.dumps(manifest), encoding="utf-8")
                dataset = read_dataset(self.input)
                self.assertEqual(len(dataset["transactions"]), 3630)
                self.assertEqual(dataset["complete_days"], [])
                self.assertTrue(any(item["code"] == "invalid_manifest_file" and "escapes" in item["message"]
                    for item in dataset["issues"]))
                # Reset the manifest before the next path variant.
                shutil.copy2(self.base / "manifest.json", self.input / "manifest.json")

    def test_source_symlink_escape_is_not_read(self):
        outside = Path(self.temp.name) / "outside.xlsx"
        shutil.copy2(self.tool2(), outside)
        try:
            (self.input / "symlink.xlsx").symlink_to(outside)
        except OSError:
            self.skipTest("This Windows account cannot create file symlinks")
        dataset = read_dataset(self.input)
        self.assertEqual(len(dataset["transactions"]), 3630)
        self.assertEqual(dataset["complete_days"], [])
        self.assertTrue(any(item["code"] == "unsafe_source_path" for item in dataset["issues"]))

    def test_fixed_seed_has_identical_native_business_contents(self):
        generated = generate_demo(Path(self.temp.name) / "second")
        for path in (Path("admin.xlsx"), Path("Tool2/2026-10-06/tool2_2026-10-06.xlsx")):
            first = load_workbook(self.input / path, read_only=True)
            second = load_workbook(generated / path, read_only=True)
            self.assertEqual(list(first.active.values), list(second.active.values))
            first.close()
            second.close()
        self.assertEqual(self.manifest(), json.loads((generated / "manifest.json").read_text(encoding="utf-8")))


class NativeFormatTests(unittest.TestCase):
    def test_identity_normalization(self):
        for raw, expected in (("00123654 - SOME.NAME", "123654"), ("123654", "123654"),
            (123654, "123654"), (12345, "012345"), ("00012345", "012345")):
            with self.subTest(raw=raw):
                self.assertEqual(normalize_agent_id(raw), expected)
        for raw in (None, "user123654", "1234567", "12345", "01123654", -1, 123.5):
            with self.subTest(raw=raw):
                self.assertRaises(ValueError, normalize_agent_id, raw)

    def test_native_duration_formats_and_ambiguous_decimal_rejection(self):
        self.assertEqual(duration_seconds(1342), 1342)
        self.assertEqual(duration_seconds("1342.0"), 1342)
        self.assertEqual(duration_seconds(time(0, 22, 22)), 1342)
        self.assertEqual(duration_seconds(timedelta(hours=27)), 97200)
        self.assertEqual(duration_seconds("1:32:34"), 5554)
        self.assertEqual(duration_seconds(time(0, 9, 47, 125000), average=True), 587.125)
        for raw in (None, -1, 21.57, "21.57", "0:99:00"):
            with self.subTest(raw=raw):
                self.assertRaises(ValueError, duration_seconds, raw)

    def test_minutes_seconds_rebuild_and_full_time_cell(self):
        base = datetime(2026, 10, 1, 18, 17)
        result = rebuild_timestamp(base, "17.16.0")
        self.assertEqual(result.second, 16)
        self.assertEqual(result.utcoffset(), timedelta(hours=-6))
        self.assertEqual(rebuild_timestamp(base, time(18, 17, 16)), result)
        self.assertRaises(ValueError, rebuild_timestamp, base, "18.16.0")
        self.assertRaises(ValueError, rebuild_timestamp, base, "17.60.0")


if __name__ == "__main__":
    unittest.main()
