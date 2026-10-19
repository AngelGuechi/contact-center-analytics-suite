"""Strict, streamed readers for the synthetic native Tool 1 / Tool 2 contract.

Source files are never changed. A workbook is admitted atomically: an invalid
business row rejects that workbook, while good workbooks on other dates survive.
Completeness is an explicit manifest assertion verified against accepted files.
"""
from __future__ import annotations

import hashlib
import json
import math
import re
from contextvars import ContextVar
from datetime import date, datetime, time, timedelta, timezone
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

PARSER_VERSION = "native-xlsx-1"
BUSINESS_ZONE = ZoneInfo("America/Mexico_City")
_RAW_CELL_POOL = ContextVar("native_xlsx_raw_cell_pool", default=None)
TOOL2_HEADERS = ["Transaction ID", "Date(Server)", "Time(Server)",
    "Date/Time(Server)", "Duration (sec)", "Connection ID", "Transferred Out",
    "Num Holds", "Login ID", "End Date/Time(Server)", "End Time(Server)",
    "POST_ROUTE_DATA"]
TOOL1_HEADERS = [None, "AgentStatus", *[f"X{i:02d}" for i in range(14)],
    "Inbound", "X14", "X15", "Transfers Made", "TalkTime", "HoldTime", "X16",
    "HandleTime", "AvgTalkTime", "AvgHoldTime", "X17", "AvgHandleTime"]
ADMIN_HEADERS = ["Agent Name", "Supervisor", "Manager", "AgentClientNumberID",
    "AgentClientUserID", "00+AgentClientNumberID", None, "POST_ROUTE_DATA",
    "CATEGORY CODE", "CATEGORY"]


class WorkbookValidationError(ValueError):
    def __init__(self, message, business_dates):
        super().__init__(message)
        self.business_dates = set(business_dates)


def normalize_agent_id(value) -> str:
    if isinstance(value, bool) or value is None:
        raise ValueError(f"Invalid Agent ID {value!r}")
    if isinstance(value, (int, float)):
        if not math.isfinite(value) or int(value) != value or value < 0:
            raise ValueError(f"Invalid Agent ID {value!r}")
        digits = str(int(value))
        if len(digits) <= 6:
            return digits.zfill(6)
    else:
        match = re.match(r"^(\d+)(?:\s|$)", str(value).strip())
        if not match:
            raise ValueError(f"Invalid Agent ID {value!r}")
        digits = match.group(1)
    if len(digits) == 8 and digits.startswith("00"):
        return digits[2:]
    if len(digits) == 6:
        return digits
    raise ValueError(f"Invalid Agent ID {value!r}; expected six digits or prefix 00")


def _count(value, name="count") -> int:
    if isinstance(value, bool) or value is None:
        raise ValueError(f"Invalid {name} {value!r}; expected nonnegative whole number")
    try:
        number = float(value)
    except (TypeError, ValueError):
        raise ValueError(f"Invalid {name} {value!r}; expected nonnegative whole number") from None
    if not math.isfinite(number) or number < 0 or int(number) != number:
        raise ValueError(f"Invalid {name} {value!r}; expected nonnegative whole number")
    return int(number)


def duration_seconds(value, *, average=False):
    """Time-typed cells, whole seconds and h:mm:ss; never guess decimal units."""
    if isinstance(value, timedelta):
        result = value.total_seconds()
    elif isinstance(value, time):
        result = value.hour * 3600 + value.minute * 60 + value.second + value.microsecond / 1e6
    elif isinstance(value, datetime):
        if value.year not in (1899, 1900):
            raise ValueError(f"Invalid duration {value!r}")
        result = duration_seconds(value.time(), average=True)
    elif isinstance(value, (int, float)) and not isinstance(value, bool):
        result = float(value)
    elif isinstance(value, str):
        text = value.strip()
        match = re.fullmatch(r"(\d+):(\d{2}):(\d{2})(?:\.(\d+))?", text)
        if match:
            hours, minutes, seconds = map(int, match.groups()[:3])
            if minutes >= 60 or seconds >= 60:
                raise ValueError(f"Invalid duration {value!r}")
            result = hours * 3600 + minutes * 60 + seconds
            if match.group(4):
                result += float("0." + match.group(4))
        elif re.fullmatch(r"\d+(?:\.0+)?", text):
            result = float(text)
        else:
            raise ValueError(f"Invalid duration {value!r}; use whole seconds or a time cell")
    else:
        raise ValueError(f"Invalid duration {value!r}")
    if not math.isfinite(result) or result < 0:
        raise ValueError(f"Invalid duration {value!r}")
    if average:
        return result
    # Excel serial times carry tiny conversion noise; fractional seconds other
    # than that are rejected rather than rounded into a different source value.
    if abs(result - round(result)) > 1e-6:
        raise ValueError(f"Invalid duration {value!r}; totals must be whole seconds")
    return int(round(result))


def _datetime(value) -> datetime:
    if isinstance(value, datetime):
        result = value
    elif isinstance(value, date):
        result = datetime.combine(value, time())
    elif isinstance(value, str):
        try:
            result = datetime.fromisoformat(value.strip().replace("Z", "+00:00"))
        except ValueError:
            raise ValueError(f"Invalid date/time {value!r}") from None
    else:
        raise ValueError(f"Invalid date/time {value!r}")
    if result.tzinfo is None:
        return result.replace(tzinfo=BUSINESS_ZONE)
    return result.astimezone(BUSINESS_ZONE)


def rebuild_timestamp(base, clock) -> datetime:
    minute = _datetime(base)
    if isinstance(clock, str) and (match := re.fullmatch(r"(\d{1,2})\.(\d{1,2})\.0", clock.strip())):
        mm, ss = map(int, match.groups())
        if mm >= 60 or ss >= 60 or mm != minute.minute or minute.second != 0:
            raise ValueError(f"Clock {clock!r} does not match minute timestamp {base!r}")
        return minute.replace(second=ss, microsecond=0)
    if isinstance(clock, datetime):
        clock = clock.time()
    if isinstance(clock, str):
        try:
            clock = time.fromisoformat(clock.strip())
        except ValueError:
            raise ValueError(f"Invalid server clock {clock!r}") from None
    if isinstance(clock, time):
        if (clock.hour, clock.minute) != (minute.hour, minute.minute):
            raise ValueError(f"Clock {clock!r} does not match timestamp {base!r}")
        return minute.replace(second=clock.second, microsecond=clock.microsecond)
    raise ValueError(f"Invalid server clock {clock!r}")


def _utc(value: datetime) -> str:
    return value.astimezone(timezone.utc).isoformat().replace("+00:00", "Z")


def _key(value) -> str:
    return " ".join(str(value or "").split()).casefold()


def _sha(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _raw_value(value):
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, timedelta):
        return {"days": value.days, "seconds": value.seconds, "microseconds": value.microseconds}
    return value


def _source(relative, digest, sheet, row, headers, cells):
    pool = _RAW_CELL_POOL.get()
    raw_cells = {}
    for index, (header, cell) in enumerate(zip(headers, cells)):
        value = _raw_value(cell.value)
        cell_type = type(cell.value).__name__
        number_format = cell.number_format
        # Immutable source descriptors are shared within one import only.
        # Blank filler cells and repeated clocks/IDs no longer allocate one
        # three-field dictionary for every column of every transaction.
        value_key = tuple(value.items()) if isinstance(value, dict) else value
        key = (cell_type, number_format, value_key)
        descriptor = pool.get(key) if pool is not None else None
        if descriptor is None:
            descriptor = dict(value=value, type=cell_type, number_format=number_format)
            if pool is not None:
                pool[key] = descriptor
        raw_cells[str(header or f"column_{index + 1}")] = descriptor
    return {"file": relative, "sha256": digest, "sheet": sheet, "row": row,
        "import_id": f"xlsx:{digest}", "raw_cells": raw_cells}


def _header_rows(sheet, required):
    """Find a native header after optional pasted title rows; stream remaining rows."""
    iterator = iter(sheet.iter_rows())
    for number in range(1, 21):
        cells = next(iterator, None)
        if cells is None:
            break
        headers = [cell.value for cell in cells]
        mapping = {}
        repeated = []
        for index, name in enumerate(headers):
            key = _key(name)
            if key in required:
                if key in mapping:
                    repeated.append(str(name))
                mapping[key] = index
        if all(key in mapping for key in required):
            if repeated:
                raise ValueError("Repeated required headers: " + ", ".join(repeated))
            return headers, mapping, number, iterator
    raise ValueError("Missing native headers: " + ", ".join(sorted(required)))


def _blank(cells):
    return all(cell.value is None or str(cell.value).strip() == "" for cell in cells)


def _roster(path, relative, digest, issue, summary):
    agents, queues, ids, logins = [], {}, set(), set()
    book = load_workbook(path, read_only=True, data_only=False)
    try:
        sheet = book["Admin"] if "Admin" in book.sheetnames else book.worksheets[0]
        required = {_key(name) for name in ADMIN_HEADERS if name}
        headers, mapping, header_row, rows = _header_rows(sheet, required)
        summary["skipped_rows"] += header_row - 1
        for row_number, cells in enumerate(rows, header_row + 1):
            if _blank(cells):
                summary["blank_rows"] += 1
                continue
            values = [cell.value for cell in cells]
            get = lambda name: values[mapping[_key(name)]]
            if any(get(name) is not None for name in ADMIN_HEADERS[:6]):
                try:
                    agent_id = normalize_agent_id(get("AgentClientNumberID"))
                    login = str(get("AgentClientUserID") or "").strip()
                    if not login or agent_id in ids or login in logins:
                        raise ValueError(f"Missing or duplicate roster identity {agent_id!r}/{login!r}")
                    names = [str(get(name) or "").strip() for name in ADMIN_HEADERS[:3]]
                    if not all(names):
                        raise ValueError("Roster name, supervisor and manager must be supplied")
                    if normalize_agent_id(get("00+AgentClientNumberID")) != agent_id:
                        raise ValueError("Prefixed Agent ID disagrees with AgentClientNumberID")
                    agents.append(dict(agent_id=agent_id, login_id=login, name=names[0],
                        supervisor=names[1], manager=names[2], sources=[_source(relative, digest,
                        sheet.title, row_number, headers, cells)]))
                    ids.add(agent_id)
                    logins.add(login)
                except ValueError as exc:
                    issue("error", "invalid_admin_row", str(exc), relative, row_number)
                    raise
            if any(get(name) is not None for name in ADMIN_HEADERS[7:]):
                try:
                    code = str(_count(get("POST_ROUTE_DATA"), "queue code"))
                    category, call_type = str(get("CATEGORY CODE") or "").strip(), str(get("CATEGORY") or "").strip()
                    if code in queues or not category or not call_type:
                        raise ValueError(f"Duplicate or incomplete queue mapping {code!r}")
                    queues[code] = call_type
                except ValueError as exc:
                    issue("error", "invalid_admin_row", str(exc), relative, row_number)
                    raise
    finally:
        book.close()
    return agents, queues


def _transactions(sheet, relative, digest, roster_by_login, queues, issue, summary):
    required = {_key(name) for name in TOOL2_HEADERS}
    headers, mapping, header_row, rows = _header_rows(sheet, required)
    summary["skipped_rows"] += header_row - 1
    records, days = [], set()
    for row_number, cells in enumerate(rows, header_row + 1):
        if _blank(cells):
            summary["blank_rows"] += 1
            continue
        values = [cell.value for cell in cells]
        get = lambda name: values[mapping[_key(name)]]
        try:
            # Track the row's own date before timing/count validation so even
            # an unlabelled bad workbook invalidates its affected known day.
            days.add(_datetime(get("Date/Time(Server)")).date().isoformat())
            transaction_id = str(_count(get("Transaction ID"), "Transaction ID"))
            if len(transaction_id) != 8:
                raise ValueError(f"Transaction ID must have eight digits: {get('Transaction ID')!r}")
            started = rebuild_timestamp(get("Date/Time(Server)"), get("Time(Server)"))
            ended = rebuild_timestamp(get("End Date/Time(Server)"), get("End Time(Server)"))
            day = started.date().isoformat()
            days.add(day)
            if _datetime(get("Date(Server)")).date() != started.date():
                raise ValueError("Date(Server) disagrees with start date")
            duration = duration_seconds(get("Duration (sec)"))
            if abs((ended - started).total_seconds() - duration) > 1:
                raise ValueError(f"End/start difference does not equal duration {duration}s within one second")
            connection = str(get("Connection ID") or "").strip()
            if not re.fullmatch(r"[a-zA-Z0-9]{16}", connection):
                raise ValueError(f"Invalid Connection ID {connection!r}; expected 16 letters/digits")
            login = str(get("Login ID") or "").strip()
            if not login:
                raise ValueError("Blank Login ID")
            agent = roster_by_login.get(login)
            agent_id = agent["agent_id"] if agent else "unmapped:" + login
            if not agent:
                issue("warning", "unmapped_login", f"Login {login!r} retained under Unassigned", relative, row_number)
            queue = str(_count(get("POST_ROUTE_DATA"), "POST_ROUTE_DATA"))
            call_type = queues.get(queue, "Unknown call type")
            if queue not in queues:
                issue("warning", "unmapped_queue", f"Queue {queue!r} retained as Unknown call type", relative, row_number)
            records.append(dict(transaction_id=transaction_id, connection_id=connection,
                agent_id=agent_id, login_id=login, call_type=call_type, queue_code=queue,
                business_date=day, started_at=_utc(started), ended_at=_utc(ended),
                duration_seconds=duration, transfers=_count(get("Transferred Out"), "Transferred Out"),
                holds=_count(get("Num Holds"), "Num Holds"),
                sources=[_source(relative, digest, sheet.title, row_number, headers, cells)]))
        except ValueError as exc:
            issue("error", "invalid_transaction", str(exc), relative, row_number)
            raise WorkbookValidationError(str(exc), days) from exc
    return records, days


def _snapshot(sheet, relative, digest, captured, roster_by_id, issue, summary):
    required = {_key(name) for name in TOOL1_HEADERS if name and not name.startswith("X")}
    headers, mapping, header_row, rows = _header_rows(sheet, required)
    summary["skipped_rows"] += header_row - 1
    records, identities = [], set()
    for row_number, cells in enumerate(rows, header_row + 1):
        if _blank(cells):
            summary["blank_rows"] += 1
            continue
        values = [cell.value for cell in cells]
        identity = values[0]
        if identity is None or not re.match(r"^\d", str(identity).strip()):
            summary["skipped_rows"] += 1
            continue
        get = lambda name: values[mapping[_key(name)]]
        try:
            agent_id = normalize_agent_id(identity)
            if agent_id in identities:
                raise ValueError(f"Repeated snapshot Agent ID {agent_id}")
            identities.add(agent_id)
            calls = _count(get("Inbound"), "Inbound")
            talk, hold, handle = (duration_seconds(get(name)) for name in ("TalkTime", "HoldTime", "HandleTime"))
            if abs(handle - talk - hold) > 1:
                raise ValueError(f"HandleTime {handle}s differs from TalkTime+HoldTime {talk + hold}s")
            for name, total in (("AvgTalkTime", talk), ("AvgHoldTime", hold), ("AvgHandleTime", handle)):
                expected = total / calls if calls else 0
                if not calls and total:
                    raise ValueError(f"Zero Inbound with nonzero total {total}")
                if abs(duration_seconds(get(name), average=True) - expected) > 1:
                    raise ValueError(f"{name} disagrees with total/Inbound within one second")
            if agent_id not in roster_by_id:
                issue("warning", "unmapped_agent", f"Agent {agent_id!r} retained under Unassigned", relative, row_number)
            records.append(dict(agent_id=agent_id, calls=calls, talk_seconds=talk,
                hold_seconds=hold, handle_seconds=handle,
                transfers=_count(get("Transfers Made"), "Transfers Made"),
                status=str(get("AgentStatus") or ""), sources=[_source(relative, digest,
                sheet.title, row_number, headers, cells)]))
        except ValueError as exc:
            issue("error", "invalid_snapshot", str(exc), relative, row_number)
            raise
    return dict(snapshot_id=f"snapshot:{digest}", captured_at=_utc(captured),
        business_date=captured.date().isoformat(), agents=records, issues=[], accepted=True,
        source_file=relative, source_sha256=digest, import_id=f"xlsx:{digest}",
        parser_version=PARSER_VERSION)


def _safe_path(root, relative):
    path = (root / relative).resolve()
    if not path.is_relative_to(root):
        raise ValueError(f"File path escapes input folder: {relative!r}")
    return path


def _business_key(record):
    return tuple((key, record[key]) for key in sorted(record) if key != "sources")


def _read_dataset(input_dir: Path) -> dict:
    """Read manifest + Admin + native .xlsx exports into an engine-neutral dataset."""
    root = Path(input_dir).resolve()
    issues = []
    def issue(severity, code, message, file=None, row=None):
        result = dict(severity=severity, code=code, message=message, file=file, row=row)
        issues.append(result)
        return result
    summary = dict(files_seen=0, files_accepted=0, files_rejected=0, files_missing=0,
        transaction_rows=0, transactions=0, duplicate_transactions=0,
        snapshots=0, duplicate_snapshots=0, blank_rows=0, skipped_rows=0, imports=[])
    manifest = {}
    try:
        manifest = json.loads((root / "manifest.json").read_text(encoding="utf-8"))
        if not isinstance(manifest, dict):
            raise ValueError("Manifest must be an object")
        for field in ("expected_days", "complete_days"):
            if not isinstance(manifest.get(field, []), list):
                raise ValueError(f"Manifest {field} must be a list")
            for day in manifest.get(field, []):
                date.fromisoformat(day)
        if not isinstance(manifest.get("files", []), list):
            raise ValueError("Manifest files must be a list")
    except (OSError, ValueError, TypeError) as exc:
        issue("warning", "manifest_unavailable", f"Completeness unverified: {exc}", "manifest.json")
        manifest = {}
    descriptors = {}
    manifest_paths_valid = True
    for entry in manifest.get("files", []):
        if not isinstance(entry, dict) or not isinstance(entry.get("file"), str):
            issue("error", "invalid_manifest_file", "Each manifest file needs a relative file path", "manifest.json")
            manifest_paths_valid = False
            continue
        relative = entry["file"].replace("\\", "/")
        try:
            _safe_path(root, relative)
        except ValueError as exc:
            issue("error", "invalid_manifest_file", str(exc), "manifest.json")
            manifest_paths_valid = False
            continue
        if relative in descriptors:
            issue("error", "invalid_manifest_file", f"Repeated file declaration {relative}", "manifest.json")
            manifest_paths_valid = False
            continue
        descriptors[relative] = dict(entry, file=relative)
    for path in sorted(root.rglob("*.xlsx")) if root.exists() else []:
        if path.name.startswith("~$"):
            continue
        relative = path.relative_to(root).as_posix()
        try:
            _safe_path(root, relative)
        except ValueError as exc:
            issue("error", "unsafe_source_path", str(exc), relative)
            manifest_paths_valid = False
            continue
        descriptors.setdefault(relative, dict(file=relative, required=False))
    roster, queues = [], {}
    admin_relative = str(manifest.get("admin_file", "admin.xlsx")).replace("\\", "/")
    admin_valid = False
    try:
        admin_path = _safe_path(root, admin_relative)
        roster, queues = _roster(admin_path, admin_relative, _sha(admin_path), issue, summary)
        admin_valid = True
    except Exception as exc:
        issue("error", "admin_unavailable", f"Admin rejected: {exc}", admin_relative)
    roster_by_login = {agent["login_id"]: agent for agent in roster}
    roster_by_id = {agent["agent_id"]: agent for agent in roster}
    all_transactions, snapshots = {}, []
    failed_days, imported_days, accepted_files = set(), set(), set()
    for relative, entry in sorted(descriptors.items()):
        if relative == admin_relative or entry.get("kind") == "admin":
            continue
        summary["files_seen"] += 1
        path = _safe_path(root, relative)
        declared_day = entry.get("business_date")
        captured = None
        try:
            capture_value = entry.get("captured_at")
            if capture_value:
                captured = _datetime(capture_value)
            elif match := re.search(r"tool1_(\d{4}-\d{2}-\d{2})_(\d{2})(\d{2})", path.stem, re.I):
                captured = _datetime(f"{match.group(1)}T{match.group(2)}:{match.group(3)}:00")
        except ValueError as exc:
            issue("error", "invalid_capture_time", str(exc), relative)
        if not path.is_file():
            if entry.get("required", True):
                issue("error", "missing_file", "Required source file is missing", relative)
                summary["files_missing"] += 1
                if declared_day:
                    failed_days.add(declared_day)
                if entry.get("kind") == "tool1" and captured:
                    snapshots.append(dict(snapshot_id=f"missing:{relative}", captured_at=_utc(captured),
                        business_date=captured.date().isoformat(), agents=[], accepted=False,
                        issues=[dict(code="missing_snapshot", message="Required snapshot missing")]))
            continue
        digest, kind, parsed_days = _sha(path), entry.get("kind"), set()
        before_issues = len(issues)
        book = None
        try:
            book = load_workbook(path, read_only=True, data_only=False)
            if kind == "tool1" or "LiveAgentData" in book.sheetnames:
                kind = "tool1"
                if captured is None:
                    raise ValueError("Snapshot requires capture time from manifest or native filename")
                if declared_day and declared_day != captured.date().isoformat():
                    raise ValueError("Declared day disagrees with snapshot capture date")
                sheet = book["LiveAgentData"]
                snapshot = _snapshot(sheet, relative, digest, captured, roster_by_id, issue, summary)
                same_time = [existing for existing in snapshots
                    if existing["captured_at"] == snapshot["captured_at"] and existing["accepted"]]
                if same_time:
                    previous = same_time[0]
                    previous_rows = {row["agent_id"]: row for row in previous["agents"]}
                    current_rows = {row["agent_id"]: row for row in snapshot["agents"]}
                    if previous_rows.keys() != current_rows.keys() or any(
                        _business_key(previous_rows[agent_id]) != _business_key(current_rows[agent_id])
                        for agent_id in previous_rows):
                        # Neither version can establish the correct total at
                        # this timestamp. The rejected marker breaks the chain.
                        previous["accepted"] = False
                        previous["issues"].append(dict(code="conflicting_snapshot",
                            message="Conflicting totals at the same captured time"))
                        previous["agents"] = []
                        raise ValueError("Conflicting snapshots at the same captured time")
                    for agent_id in previous_rows:
                        previous_rows[agent_id]["sources"].extend(current_rows[agent_id]["sources"])
                    summary["duplicate_snapshots"] += 1
                else:
                    snapshots.append(snapshot)
                parsed_days.add(captured.date().isoformat())
            else:
                kind = "tool2"
                sheet = book["Export"] if "Export" in book.sheetnames else book.worksheets[0]
                records, parsed_days = _transactions(sheet, relative, digest, roster_by_login, queues, issue, summary)
                if declared_day and parsed_days != {declared_day}:
                    raise ValueError("Declared day disagrees with transaction start dates, or no rows verify the export date")
                # Validate duplicates before committing any part of this file.
                local = {}
                for record in records:
                    prior = local.get(record["transaction_id"]) or all_transactions.get(record["transaction_id"])
                    if prior and _business_key(prior) != _business_key(record):
                        raise ValueError(f"Conflicting duplicate Transaction ID {record['transaction_id']}; workbook rejected")
                    local.setdefault(record["transaction_id"], record)
                for record in records:
                    transaction_id = record["transaction_id"]
                    if transaction_id in all_transactions:
                        all_transactions[transaction_id]["sources"].extend(record["sources"])
                        summary["duplicate_transactions"] += 1
                    else:
                        all_transactions[transaction_id] = record
                summary["transaction_rows"] += len(records)
                imported_days.update(parsed_days)
            summary["files_accepted"] += 1
            accepted_files.add(relative)
            summary["imports"].append(dict(file=relative, sha256=digest, import_id=f"xlsx:{digest}",
                kind=kind, accepted=True, business_dates=sorted(parsed_days), parser_version=PARSER_VERSION))
        except Exception as exc:
            issue("error", "rejected_file", str(exc), relative)
            summary["files_rejected"] += 1
            parsed_days.update(getattr(exc, "business_dates", set()))
            failed_days.update(parsed_days)
            if declared_day:
                failed_days.add(declared_day)
            elif match := re.search(r"\d{4}-\d{2}-\d{2}", path.name):
                # A name can show which declared day an error affects, but it
                # cannot establish that any export is complete.
                failed_days.add(match.group())
            if kind == "tool1" and captured:
                snapshots.append(dict(snapshot_id=f"rejected:{digest}", captured_at=_utc(captured),
                    business_date=captured.date().isoformat(), agents=[], accepted=False,
                    issues=issues[before_issues:], source_file=relative, source_sha256=digest,
                    parser_version=PARSER_VERSION))
            summary["imports"].append(dict(file=relative, sha256=digest, import_id=f"xlsx:{digest}",
                kind=kind, accepted=False, business_dates=sorted(parsed_days), parser_version=PARSER_VERSION))
        finally:
            if book is not None:
                book.close()
    for record in all_transactions.values():
        agent_id = record["agent_id"]
        if agent_id not in roster_by_id:
            agent = dict(agent_id=agent_id, login_id=record["login_id"], name="Unassigned",
                supervisor="Unassigned", manager="Unassigned")
            roster.append(agent)
            roster_by_id[agent_id] = agent
    for snapshot in snapshots:
        for row in snapshot["agents"]:
            if row["agent_id"] not in roster_by_id:
                agent = dict(agent_id=row["agent_id"], login_id="", name="Unassigned",
                    supervisor="Unassigned", manager="Unassigned")
                roster.append(agent)
                roster_by_id[agent["agent_id"]] = agent
    complete = []
    for day in sorted(set(manifest.get("complete_days", []))):
        required = [entry for entry in descriptors.values() if entry.get("business_date") == day
            and entry.get("kind") == "tool2" and entry.get("required", True)]
        matched = [entry for entry in required if entry.get("complete", False)]
        if manifest_paths_valid and admin_valid and day in manifest.get("expected_days", []) and day in imported_days and day not in failed_days and matched and all(
            entry["file"] in accepted_files for entry in required):
            complete.append(day)
        else:
            issue("warning", "incomplete_day", f"Day {day} is not verified complete", "manifest.json")
    transactions = sorted(all_transactions.values(), key=lambda row: (row["business_date"], row["transaction_id"]))
    snapshots.sort(key=lambda row: (row["captured_at"], row["snapshot_id"]))
    summary["transactions"] = len(transactions)
    summary["snapshots"] = sum(snapshot["accepted"] for snapshot in snapshots)
    return dict(roster=roster, transactions=transactions, snapshots=snapshots,
        complete_days=complete, expected_days=sorted(set(manifest.get("expected_days", []))),
        manifest=manifest, issues=issues, import_summary=summary, parser_version=PARSER_VERSION)


def read_dataset(input_dir: Path) -> dict:
    """Read a dataset with an isolated descriptor pool, safe across concurrent imports."""
    token = _RAW_CELL_POOL.set({})
    try:
        return _read_dataset(input_dir)
    finally:
        _RAW_CELL_POOL.reset(token)
