"""Generate use-case-owned, clearly synthetic native Excel input fixtures.

These workbooks exercise the actual importer; they are not a prepared report.
No original AHT fixtures or operational data are read or overwritten.
"""
from __future__ import annotations

import json
import random
from datetime import date, datetime, time, timedelta
from pathlib import Path

from openpyxl import Workbook

from .readers import ADMIN_HEADERS, TOOL1_HEADERS, TOOL2_HEADERS

AGENTS = [
    ("123654", "SYNANA", "Synthetic Ana Ruiz", "Synthetic Supervisor One", "Synthetic Manager"),
    ("235694", "SYNBRUNO", "Synthetic Bruno Mora", "Synthetic Supervisor One", "Synthetic Manager"),
    ("345678", "SYNCARLA", "Synthetic Carla Lopez", "Synthetic Supervisor Two", "Synthetic Manager"),
    ("456789", "SYNDIEGO", "Synthetic Diego Perez", "Synthetic Supervisor Two", "Synthetic Manager"),
    ("012345", "SYNEVA", "Synthetic Eva Soto", "Synthetic Supervisor Two", "Synthetic Manager"),
    ("567890", "SYNFINN", "Synthetic Finn Zero", "Synthetic Supervisor Two", "Synthetic Manager"),
]
QUEUES = [(8167, "AS_SDFE_OSME", "MIMO"), (99920, "sa_WRW_md", "NEM"),
    (8112, "op_aasf_RTYT", "Credit"), (8125, "pl_AASD_ER", "General Billing"),
    (8140, "lp_KQW_tr", "LPC"), (8153, "CB_ZXRP_mn", "CBB")]


def _workbook(sheet_name):
    book = Workbook()
    book.properties.creator = "Call Avoidance Detector synthetic generator"
    book.properties.title = "SYNTHETIC DEMO DATA — not operational agent data"
    # Contents, including metadata, are stable for the same arguments.
    book.properties.created = datetime(2026, 1, 1)
    book.properties.modified = datetime(2026, 1, 1)
    book.active.title = sheet_name
    return book


def _save(book, path):
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)
    book.close()


def _write_admin(destination):
    book = _workbook("Admin")
    sheet = book.active
    sheet.append(ADMIN_HEADERS)
    for index in range(max(len(AGENTS), len(QUEUES))):
        roster = [None] * 6
        if index < len(AGENTS):
            agent_id, login, name, supervisor, manager = AGENTS[index]
            roster = [name, supervisor, manager, int(agent_id), login, "00" + agent_id]
        queue = list(QUEUES[index]) if index < len(QUEUES) else [None] * 3
        sheet.append(roster + [None] + queue)
    _save(book, destination / "admin.xlsx")


def _write_tool2(path, records):
    book = _workbook("Export")
    sheet = book.active
    # Native business headers plus thirty hidden fillers = exactly 42 columns.
    headers = TOOL2_HEADERS + [f"X{i:02d}" for i in range(1, 31)]
    sheet.append(headers)
    for index, record in enumerate(records):
        started, ended = record["started"], record["ended"]
        # Cover both permitted native Duration formats in every exported day.
        duration = record["duration"] if index % 2 else timedelta(seconds=record["duration"])
        start_clock = f"{started.minute:02d}.{started.second:02d}.0"
        end_clock = f"{ended.minute:02d}.{ended.second:02d}.0"
        if index % 11 == 0:
            start_clock, end_clock = started.time(), ended.time()
        sheet.append([record["transaction_id"], datetime.combine(started.date(), time()),
            start_clock, started.replace(second=0, microsecond=0), duration,
            record["connection_id"], record["transfers"], record["holds"], record["login"],
            ended.replace(second=0, microsecond=0), end_clock, record["queue"], *[None] * 30])
        row = sheet.max_row
        for col in (2, 4, 10):
            sheet.cell(row, col).number_format = "yyyy-mm-dd hh:mm"
        if isinstance(duration, timedelta):
            sheet.cell(row, 5).number_format = "[h]:mm:ss"
        for col in (3, 11):
            if isinstance(sheet.cell(row, col).value, time):
                sheet.cell(row, col).number_format = "hh:mm:ss"
    # An explicit blank row tests counted skips without changing metrics.
    sheet.append([None] * 42)
    for col in range(13, 43):
        sheet.column_dimensions[sheet.cell(1, col).column_letter].hidden = True
    _save(book, path)


def _write_tool1(path, captured, records):
    book = _workbook("LiveAgentData")
    sheet = book.active
    sheet.append(TOOL1_HEADERS)
    for agent_id, login, name, _, _ in AGENTS:
        completed = [record for record in records if record["agent_id"] == agent_id and record["ended"] <= captured]
        calls = len(completed)
        talk = sum(record["talk"] for record in completed)
        hold = sum(record["hold"] for record in completed)
        handle = talk + hold
        transfers = sum(record["transfers"] for record in completed)
        row = ["00" + agent_id + " - " + name, "ReadyForNextCall (00:01:00)", *["Y"] * 14,
            calls, None, None, transfers, timedelta(seconds=talk), timedelta(seconds=hold),
            None, timedelta(seconds=handle), timedelta(seconds=talk / calls if calls else 0),
            timedelta(seconds=hold / calls if calls else 0), None,
            timedelta(seconds=handle / calls if calls else 0)]
        sheet.append(row)
        for col in (21, 22, 24, 25, 26, 28):
            sheet.cell(sheet.max_row, col).number_format = "[h]:mm:ss.000"
    sheet.append(["TOTAL (synthetic title row)"])
    incoming = book.create_sheet("LiveIncomingData")
    incoming.append([None, "Current Calls Waiting", "Current Max Waiting", "X01", "X02", "X03",
        "Entered", "Answered", "Abandoned", *[f"X{i:02d}" for i in range(7, 12)]])
    for queue, category, _ in QUEUES:
        incoming.append([category, 0, timedelta(), None, None, None, 100, 80, 10, *[None] * 5])
        incoming.cell(incoming.max_row, 3).number_format = "[h]:mm:ss"
    _save(book, path)


def generate_demo(destination: Path, seed: int = 5, report_date: str = "2026-10-06") -> Path:
    """Create 15 declared complete native Tool 2 dates plus today's Tool 1 replay.

    Existing non-generator files survive. Regenerating replaces only the named
    synthetic fixture outputs; do not point this function at operational inputs.
    """
    destination = Path(destination).resolve()
    destination.mkdir(parents=True, exist_ok=True)
    report = date.fromisoformat(report_date)
    rng = random.Random(seed)
    days = [report - timedelta(days=offset) for offset in range(14, -1, -1)]
    _write_admin(destination)
    manifest = dict(schema_version=1, synthetic=True, generator="case5-native-fixtures-1",
        seed=seed, report_date=report_date, server_timezone="America/Mexico_City",
        admin_file="admin.xlsx", expected_days=[day.isoformat() for day in days],
        complete_days=[day.isoformat() for day in days], files=[],
        annotations=[dict(annotation_id="synthetic-incident-1", site="Synthetic single site",
            started_at=report_date + "T09:00:00-06:00", ended_at=report_date + "T10:00:00-06:00",
            note="Synthetic telephony issue reported during the weak-lead interval; human review is required.")])
    transaction_id, strong_signals, transfer_context = 70000000, [], []
    for day in days:
        records = []
        for agent_index, (agent_id, login, _, _, _) in enumerate(AGENTS):
            current = day == report
            count = (30 if agent_index in (0, 1, 2, 4) else 10) if current else 50
            if agent_index == 5:
                count = 0
            cursor = datetime.combine(day, time(8, 0))
            for index in range(count):
                transaction_id += 1
                transfers = int(index in (8, 18))
                signal = index == 0
                if current and agent_index == 0:
                    signal = index < 6
                if current and agent_index == 3:
                    signal = index < 3
                if current and agent_index == 4:
                    # Two signals, plus a short transferred transaction retained in N.
                    signal = index < 3
                    transfers = int(index == 2 or index in (8, 18))
                duration = (10 if signal else rng.randint(220, 430))
                hold = 0 if signal else rng.randint(4, 18)
                talk = duration - hold
                started = cursor
                if current and agent_index == 1:
                    transfers = 0
                    if index < 10:
                        talk, hold = 300, 6
                        started = datetime.combine(day, time(8, 0)) + timedelta(seconds=index * 309)
                    elif index < 15:
                        talk, hold = 24, 2
                        started = datetime.combine(day, time(9, 5)) + timedelta(seconds=(index - 10) * 29)
                    elif index < 20:
                        talk, hold = 30, 2
                        started = datetime.combine(day, time(9, 35)) + timedelta(seconds=(index - 15) * 35)
                    else:
                        talk, hold = 290, 10
                        started = datetime.combine(day, time(11, 0)) + timedelta(seconds=(index - 20) * 303)
                    duration = talk + hold
                ended = started + timedelta(seconds=duration)
                cursor = ended + timedelta(seconds=4)
                connection = f"SYN{transaction_id:013d}"
                # Plant one real transfer chain sharing Connection ID without merging transactions.
                if agent_index == 2 and index == 8:
                    connection = f"XFR{int(day.strftime('%Y%m%d')):013d}"
                if agent_index == 4 and index == 9:
                    connection = f"XFR{int(day.strftime('%Y%m%d')):013d}"
                record = dict(transaction_id=transaction_id, connection_id=connection,
                    agent_id=agent_id, login=login, queue=8125, started=started,
                    ended=ended, duration=duration, talk=talk, hold=hold,
                    transfers=transfers, holds=int(hold > 0))
                records.append(record)
                if current and agent_index == 0 and signal and not transfers:
                    strong_signals.append(str(transaction_id))
                if current and agent_index == 4 and signal and transfers:
                    transfer_context.append(str(transaction_id))
        relative = f"Tool2/{day.isoformat()}/tool2_{day.isoformat()}.xlsx"
        _write_tool2(destination / relative, records)
        manifest["files"].append(dict(file=relative, kind="tool2", business_date=day.isoformat(),
            required=True, complete=True))
        if day == report:
            for hour, minute in ((9, 0), (9, 30), (10, 0), (16, 0)):
                captured = datetime.combine(day, time(hour, minute))
                relative = f"Tool1/{day.isoformat()}/tool1_{day.isoformat()}_{hour:02d}{minute:02d}.xlsx"
                _write_tool1(destination / relative, captured, records)
                manifest["files"].append(dict(file=relative, kind="tool1", business_date=day.isoformat(),
                    captured_at=captured.isoformat() + "-06:00", required=True))
    answer_key = dict(synthetic=True, report_date=report_date,
        strong_daily=dict(agent_id="123654", name="Synthetic Ana Ruiz", call_type="General Billing",
            signals=6, transactions=30, rate=0.2, signal_transaction_ids=strong_signals,
            self_signals=14, self_transactions=700, self_rate=0.02,
            peer_signals=56, peer_transactions=2800, peer_rate=0.02,
            reasons=["peer_outlier", "self_change"]),
        weak_intraday=dict(agent_id="235694", name="Synthetic Bruno Mora", calls=[10, 15, 20],
            talk_seconds=[3000, 3120, 3270], local_capture_times=["09:00", "09:30", "10:00"],
            interval_mean_talk_seconds=[24, 30], expected_leads=1),
        low_sample=dict(agent_id="456789", signals=3, transactions=10, flag=False),
        transfer_exclusion=dict(agent_id="012345", signals=2, transactions=30,
            transferred_short_transaction_ids=transfer_context, flag=False),
        normal=dict(agent_id="345678", signals=1, transactions=30, flag=False),
        zero_calls=dict(agent_id="567890", transactions=0, rate=None, flag=False),
        source_totals=dict(transactions=14 * 250 + 130, complete_days=15, snapshots=4))
    (destination / "manifest.json").write_text(json.dumps(manifest, indent=2) + "\n", encoding="utf-8")
    (destination / "answer_key.json").write_text(json.dumps(answer_key, indent=2) + "\n", encoding="utf-8")
    return destination
